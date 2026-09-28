"""Nạp mật khẩu ĐÃ BIẾT cho tài khoản tạo trước khi có bản mã hoá hai chiều.

bcrypt một chiều nên không lấy lại được mật khẩu cũ; nhưng có danh sách (tên đăng nhập, mật khẩu)
từ nơi khác thì so được với hash. Chỉ dòng KHỚP hash mới được lưu `password_enc` — không đổi
mật khẩu của ai, không ghi gì từ dòng sai.

Hai bước như nhập tài khoản: ``apply=False`` chỉ kiểm, ``apply=True`` kiểm lại rồi mới lưu.
Mật khẩu không bao giờ nằm trong kết quả trả về.
"""
import asyncio
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.core.errors import AppError
from app.core.security import verify_password
from app.db.mongo import get_db
from app.users import password_vault
from app.users.import_excel import (
    HASH_WORKERS,
    HEADER_SCAN_ROWS,
    MAX_ROWS,
    _field_for_header,
    _text,
    _text_password,
)
from app.users.roles import NOT_DELETED, SUPER_ADMIN_ROLE

TEMPLATE_HEADERS = ["Tên đăng nhập", "Mật khẩu"]


def _find_header(rows: list[tuple]) -> tuple[int, int, int]:
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        columns: dict[str, int] = {}
        for col, cell in enumerate(row):
            field = _field_for_header(cell)
            if field in ("username", "password") and field not in columns:
                columns[field] = col
        if "username" in columns and "password" in columns:
            return index, columns["username"], columns["password"]
    raise AppError(
        "IMPORT_NO_HEADER",
        "Không thấy dòng tiêu đề có cột Tên đăng nhập và Mật khẩu trong 5 dòng đầu.",
        422,
    )


def parse_workbook(content: bytes) -> list[dict]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — openpyxl ném đủ loại lỗi cho file hỏng/không phải xlsx
        raise AppError("IMPORT_BAD_FILE", "Không đọc được file. Chỉ nhận file Excel .xlsx.", 422)
    try:
        rows = [tuple(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()

    header_index, user_col, pass_col = _find_header(rows)
    out: list[dict] = []
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        raw_user = row[user_col] if user_col < len(row) else None
        raw_pass = row[pass_col] if pass_col < len(row) else None
        username = _text(raw_user).lower()
        password = "" if raw_pass is None else _text_password(raw_pass)
        if not username and not password:
            continue
        out.append({"row": offset, "username": username, "password": password})
        if len(out) > MAX_ROWS:
            raise AppError("IMPORT_TOO_MANY_ROWS", f"File quá {MAX_ROWS} dòng.", 422)
    if not out:
        raise AppError("IMPORT_EMPTY", "File không có dòng nào dưới dòng tiêu đề.", 422)
    return out


async def _accounts(usernames: list[str]) -> dict[str, dict]:
    if not usernames:
        return {}
    cursor = get_db().users.find(
        {"username": {"$in": usernames}, "role": {"$ne": SUPER_ADMIN_ROLE}, **NOT_DELETED},
        {"username": 1, "password_hash": 1, "password_enc": 1},
    )
    return {doc["username"]: doc async for doc in cursor}


def _summary(rows: list[dict]) -> dict:
    counts: dict[str, int] = {}
    for row in rows:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    return counts


async def run(content: bytes, *, apply: bool) -> dict:
    if apply and not password_vault.enabled():
        raise AppError(
            "PASSWORD_VAULT_DISABLED",
            "Máy chủ chưa cấu hình khoá lưu mật khẩu (PASSWORD_VAULT_KEY).",
            503,
        )
    raws = parse_workbook(content)
    rows: list[dict] = []
    first_seen: dict[str, int] = {}
    for raw in raws:
        row = {"row": raw["row"], "username": raw["username"], "status": "", "message": ""}
        if not raw["username"]:
            row.update(status="error", message="Thiếu tên đăng nhập.")
        elif not raw["password"]:
            row.update(status="error", message="Thiếu mật khẩu.")
        elif raw["username"] in first_seen:
            row.update(
                status="duplicate_in_file",
                message=f"Trùng tên đăng nhập với dòng {first_seen[raw['username']]}.",
            )
        else:
            first_seen[raw["username"]] = raw["row"]
        rows.append(row)

    accounts = await _accounts(list(first_seen))
    passwords = {raw["row"]: raw["password"] for raw in raws}
    gate = asyncio.Semaphore(HASH_WORKERS)

    async def check(row: dict) -> None:
        account = accounts.get(row["username"])
        if account is None:
            row.update(status="not_found", message="Không có tài khoản này (hoặc đã xóa).")
            return
        password = passwords[row["row"]]
        async with gate:
            ok = await asyncio.to_thread(verify_password, password, account.get("password_hash") or "")
        if not ok:
            row.update(status="mismatch", message="Mật khẩu không khớp với mật khẩu đang dùng.")
        elif password_vault.decrypt(account.get("password_enc")) == password:
            row.update(status="already", message="Đã lưu từ trước.")
        else:
            row.update(status="store", message="Khớp — sẽ lưu.")

    await asyncio.gather(*(check(row) for row in rows if not row["status"]))

    if apply:
        users = get_db().users
        for row in rows:
            if row["status"] != "store":
                continue
            account = accounts[row["username"]]
            enc_set, _ = password_vault.enc_fields(passwords[row["row"]])
            # Khớp cả hash cũ: mật khẩu vừa bị đổi giữa lúc kiểm và lúc ghi thì không lưu nhầm.
            res = await users.update_one(
                {"_id": account["_id"], "password_hash": account.get("password_hash")},
                {"$set": enc_set},
            )
            if res.matched_count:
                row.update(status="stored", message="Đã lưu.")
            else:
                row.update(status="mismatch", message="Mật khẩu vừa bị đổi, không lưu.")

    return {"applied": apply, "summary": _summary(rows), "rows": rows}


def template_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mật khẩu"
    sheet.append(TEMPLATE_HEADERS)
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 22
    # Cột mật khẩu dạng chữ: mật khẩu toàn số không bị Excel đổi thành số/mất số 0 đầu.
    for row in range(2, MAX_ROWS + 2):
        sheet.cell(row=row, column=2).number_format = "@"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
