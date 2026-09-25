"""Nhập tài khoản hàng loạt từ file Excel (trang Quản lý tài khoản, chỉ admin).

Hai bước, cùng một hàm: ``apply=False`` chỉ kiểm từng dòng và trả bảng xem trước, KHÔNG ghi gì;
``apply=True`` kiểm lại từ đầu rồi mới tạo. Không giữ file hay mật khẩu giữa hai bước — FE gửi
lại đúng file khi xác nhận, nên dữ liệu trong DB đổi giữa hai lần bấm cũng được tính lại.

Username đã có trong DB (kể cả tài khoản đã xóa mềm — unique index vẫn giữ tên) thì BỎ QUA,
không đụng gì tới tài khoản cũ. Mật khẩu KHÔNG bao giờ nằm trong kết quả trả về.
"""
import asyncio
import unicodedata
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from pymongo.errors import DuplicateKeyError

from app.core.errors import AppError
from app.core.security import hash_password
from app.db.mongo import get_db
from app.locations.catalog import ward_name_matches
from app.users.schemas import UserCreate
from app.users.service import _now


def fold(value: str) -> str:
    """Bỏ dấu + chữ thường để khớp TIÊU ĐỀ cột và nhãn Role (không dùng cho tên xã — xã so có
    dấu trước, xem catalog.ward_name_matches)."""
    value = str(value or "").replace("Đ", "D").replace("đ", "d")
    return "".join(
        char for char in unicodedata.normalize("NFD", value)
        if unicodedata.category(char) != "Mn"
    ).casefold().strip()

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
# Dòng tiêu đề có thể nằm dưới vài dòng tiêu đề trang/ghi chú.
HEADER_SCAN_ROWS = 5
# bcrypt rounds=12 tốn ~170ms/mật khẩu và chạy ĐỒNG BỘ: băm thẳng trong route async là treo cả
# server hàng chục giây với file ~100 dòng. Đẩy sang thread (bcrypt nhả GIL) và giới hạn luồng.
HASH_WORKERS = 4

ROLE_LABELS = {"province": "Hành chính công tỉnh", "commune": "Hành chính công xã"}
# File nhập CHỈ tạo được hai loại tài khoản nghiệp vụ. Không có đường nào để một file tải lên
# tạo admin / tài khoản xem thống kê tỉnh.
_ROLE_BY_TEXT = {
    fold(label): role for role, label in ROLE_LABELS.items()
} | {"province": "province", "commune": "commune"}

# Tên cột (đã bỏ dấu) → trường. Khớp BẰNG hoặc là TIỀN TỐ của tiêu đề; nhiều cột cùng khớp thì
# lấy bí danh dài nhất ("ten tinh" thắng "ten"). Cột lạ (STT, Mã cơ quan, Ghi chú…) bỏ qua.
_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("ten", "ten don vi", "ho ten", "ten hien thi"),
    "tinh": ("ten tinh", "tinh", "tinh/thanh pho", "tinh/thanh", "tinh thanh"),
    "xa": ("ten xa/phuong", "ten xa", "ten phuong/xa", "xa/phuong", "phuong/xa", "xa", "phuong"),
    "username": ("username", "ten dang nhap", "tai khoan", "tai khoan dang nhap"),
    "password": ("password", "mat khau"),
    "role": ("role", "vai tro", "loai tai khoan"),
}
_REQUIRED = ("tinh", "username", "password", "role")
_HEADER_LABELS = {
    "tinh": "Tên tỉnh", "username": "Tên đăng nhập", "password": "Mật khẩu", "role": "Role",
}
TEMPLATE_HEADERS = ["STT", "Tên", "Tên tỉnh", "Tên xã/phường", "Tên đăng nhập", "Mật khẩu", "Role"]


def _text(value) -> str:
    """Ô Excel → chuỗi. Ô số nguyên openpyxl trả float/int: 12345678 phải ra "12345678"."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _field_for_header(header: str) -> str | None:
    text = fold(" ".join(str(header or "").split()))
    if not text:
        return None
    best, best_len = None, 0
    for field, aliases in _HEADER_ALIASES.items():
        for alias in aliases:
            if (text == alias or text.startswith(alias + " ")) and len(alias) > best_len:
                best, best_len = field, len(alias)
    return best


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        columns: dict[str, int] = {}
        for col, cell in enumerate(row):
            field = _field_for_header(cell)
            if field and field not in columns:
                columns[field] = col
        if "username" in columns and "password" in columns:
            missing = [f for f in _REQUIRED if f not in columns]
            if missing:
                raise AppError(
                    "IMPORT_MISSING_COLUMNS",
                    "File thiếu cột: " + ", ".join(_HEADER_LABELS[f] for f in missing) + ".",
                    422,
                )
            return index, columns
    raise AppError(
        "IMPORT_NO_HEADER",
        "Không thấy dòng tiêu đề có cột Tên đăng nhập và Mật khẩu trong 5 dòng đầu. "
        "Hãy dùng file mẫu.",
        422,
    )


def parse_workbook(content: bytes) -> list[dict]:
    """Đọc sheet ĐẦU TIÊN → danh sách dòng thô (kèm số dòng Excel để báo lỗi đúng chỗ)."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — openpyxl ném đủ loại lỗi cho file hỏng/không phải xlsx
        raise AppError("IMPORT_BAD_FILE", "Không đọc được file. Chỉ nhận file Excel .xlsx.", 422)
    try:
        sheet = workbook.worksheets[0]
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    header_index, columns = _find_header(rows)
    out: list[dict] = []
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        cell = lambda field: row[columns[field]] if field in columns and columns[field] < len(row) else None  # noqa: E731
        values = {field: _text(cell(field)) for field in _HEADER_ALIASES}
        # Mật khẩu: giữ khoảng trắng GIỮA chuỗi (có thể là chủ ý), chỉ bỏ hai đầu.
        raw_password = cell("password")
        values["password"] = "" if raw_password is None else _text_password(raw_password)
        if not any(values[f] for f in ("username", "password", "tinh", "xa", "name")):
            continue  # dòng trống / dòng kẻ trang trí
        out.append({"row": offset, **values})
        if len(out) > MAX_ROWS:
            raise AppError("IMPORT_TOO_MANY_ROWS", f"File quá {MAX_ROWS} dòng.", 422)
    if not out:
        raise AppError("IMPORT_EMPTY", "File không có dòng tài khoản nào dưới dòng tiêu đề.", 422)
    return out


def _text_password(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _validate(raw: dict) -> dict:
    """Kiểm MỘT dòng. Trả dòng kết quả (không kèm mật khẩu); lỗi thì status "error"."""
    result = {
        "row": raw["row"], "name": raw["name"], "username": raw["username"].lower(),
        "tinh": raw["tinh"], "xa": raw["xa"], "role": "", "status": "create", "message": "",
    }

    def error(message: str) -> dict:
        return {**result, "status": "error", "message": message}

    role = _ROLE_BY_TEXT.get(fold(raw["role"]))
    if not role:
        return error("Role chỉ nhận 'Hành chính công tỉnh' hoặc 'Hành chính công xã'.")
    result["role"] = role

    try:
        created = UserCreate(username=raw["username"], password=raw["password"],
                             name=raw["name"] or None, role=role)
    except ValidationError as exc:
        return error(str(exc.errors()[0]["msg"]).removeprefix("Value error, "))
    result["username"] = created.username

    province, wards = ward_name_matches(raw["tinh"], raw["xa"])
    if not province:
        return error(f"Tỉnh '{raw['tinh']}' không có trong danh mục hiện hành.")
    result["tinh"] = province
    if raw["xa"]:
        if not wards:
            return error(f"Xã/phường '{raw['xa']}' không thuộc {province}.")
        if len(wards) > 1:
            return error(f"'{raw['xa']}' khớp nhiều đơn vị ({', '.join(wards)}) — ghi đúng dấu.")
        result["xa"] = wards[0]
    elif role == "commune":
        return error("Tài khoản Hành chính công xã phải có tên xã/phường.")
    return result


async def _existing_usernames(usernames: list[str]) -> dict[str, bool]:
    """username → đã xóa mềm hay chưa. KHÔNG lọc NOT_DELETED: unique index giữ cả tên đã xóa."""
    if not usernames:
        return {}
    cursor = get_db().users.find(
        {"username": {"$in": usernames}}, {"username": 1, "deleted_at": 1},
    )
    return {doc["username"]: doc.get("deleted_at") is not None async for doc in cursor}


def _summary(rows: list[dict]) -> dict:
    counts: dict[str, int] = {}
    for row in rows:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    return counts


async def run(content: bytes, *, apply: bool) -> dict:
    raws = parse_workbook(content)
    rows = [_validate(raw) for raw in raws]
    passwords = {raw["row"]: raw["password"] for raw in raws}

    first_seen: dict[str, int] = {}
    for row in rows:
        if row["status"] != "create":
            continue
        if row["username"] in first_seen:
            row["status"] = "duplicate_in_file"
            row["message"] = f"Trùng tên đăng nhập với dòng {first_seen[row['username']]}."
        else:
            first_seen[row["username"]] = row["row"]

    existing = await _existing_usernames(list(first_seen))
    for row in rows:
        if row["status"] == "create" and row["username"] in existing:
            deleted = existing[row["username"]]
            row["status"] = "exists_deleted" if deleted else "exists"
            row["message"] = (
                "Đã tồn tại (đã xóa — khôi phục ở mục Đã xóa)." if deleted else "Đã tồn tại, bỏ qua."
            )

    if apply:
        await _create_rows([row for row in rows if row["status"] == "create"], passwords)

    return {"applied": apply, "summary": _summary(rows), "rows": rows}


async def _create_rows(rows: list[dict], passwords: dict[int, str]) -> None:
    gate = asyncio.Semaphore(HASH_WORKERS)

    async def hashed(row: dict) -> str:
        async with gate:
            return await asyncio.to_thread(hash_password, passwords[row["row"]])

    hashes = await asyncio.gather(*(hashed(row) for row in rows))
    users = get_db().users
    for row, password_hash in zip(rows, hashes):
        now = _now()
        try:
            await users.insert_one({
                "username": row["username"],
                "password_hash": password_hash,
                "name": row["name"] or None,
                "xa": row["xa"] or None,
                "tinh": row["tinh"],
                "role": row["role"],
                "access_disabled": False,
                "created_at": now,
                "updated_at": now,
            })
        except DuplicateKeyError:
            # Ai đó vừa tạo đúng tên này giữa lúc kiểm và lúc ghi.
            row["status"] = "exists"
            row["message"] = "Đã tồn tại, bỏ qua."
            continue
        row["status"] = "created"
        row["message"] = ""


def template_bytes() -> bytes:
    """File mẫu: sheet đầu CHỈ có tiêu đề (dòng ví dụ dễ bị nhập nhầm thành tài khoản thật),
    cột Role có danh sách thả xuống; hướng dẫn nằm ở sheet thứ hai."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tài khoản"
    sheet.append(TEMPLATE_HEADERS)
    for column, width in zip("ABCDEFG", (6, 30, 16, 22, 22, 18, 24)):
        sheet.column_dimensions[column].width = width
    roles = DataValidation(
        type="list", formula1='"' + ",".join(ROLE_LABELS.values()) + '"', allow_blank=False,
    )
    sheet.add_data_validation(roles)
    roles.add(f"G2:G{MAX_ROWS + 1}")

    guide = workbook.create_sheet("Hướng dẫn")
    for line in (
        "Chỉ điền ở sheet 'Tài khoản'. Mỗi dòng một tài khoản.",
        "Tên tỉnh: tên trần hoặc đầy đủ (Bắc Ninh / Tỉnh Bắc Ninh).",
        "Tên xã/phường: CHỈ tên, không cần tiền tố Xã/Phường (vd Bồng Lai). Ghi đúng dấu.",
        "Tài khoản Hành chính công tỉnh: để trống Tên xã/phường.",
        "Tên đăng nhập: từ 3 ký tự, không phân biệt hoa thường. Đã có thì bỏ qua, không ghi đè.",
        "Mật khẩu: từ 8 ký tự.",
        "Role: Hành chính công tỉnh hoặc Hành chính công xã.",
    ):
        guide.append([line])
    guide.column_dimensions["A"].width = 90

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
