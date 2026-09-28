"""Xuất danh sách tài khoản Hành chính công xã/tỉnh ra Excel (kèm mật khẩu nếu có bản lưu).

Mỗi tỉnh/thành một sheet (tên sheet = tên tỉnh) để bàn giao riêng cho từng tỉnh.

Chỉ role commune/province, chưa xóa, đã gán tỉnh. Tỉnh nhóm theo tên CHUẨN HOÁ để "Đà Nẵng" và
"Thành phố Đà Nẵng" không tách thành hai. Mật khẩu chỉ bcrypt (chưa có bản mã hoá) → ô trống.
"""
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.core.errors import AppError
from app.db.mongo import get_db
from app.reports.integration import fold, province_name
from app.users import password_vault
from app.users.roles import NOT_DELETED, OFFICIAL_ACCOUNT_ROLES

ROLE_LABELS = {"commune": "Hành chính công xã", "province": "Hành chính công tỉnh"}
HEADERS = ["Tên đăng nhập", "Mật khẩu", "Tên", "Tỉnh", "Xã/Phường", "Vai trò", "Trạng thái"]
_WIDTHS = (26, 18, 36, 24, 28, 22, 14)
_VN = timezone(timedelta(hours=7))


async def _accounts() -> list[dict]:
    cursor = get_db().users.find(
        {**NOT_DELETED, "role": {"$in": sorted(OFFICIAL_ACCOUNT_ROLES)}},
        {"username": 1, "name": 1, "tinh": 1, "xa": 1, "role": 1, "access_disabled": 1,
         "password_enc": 1},
    )
    out = []
    async for doc in cursor:
        label = province_name(doc.get("tinh"))
        if not label:
            continue  # chưa gán tỉnh → không thuộc tỉnh nào để xuất
        out.append({**doc, "_province": label, "_province_key": fold(label)})
    return out


async def options() -> dict:
    provinces: dict[str, dict] = {}
    for doc in await _accounts():
        item = provinces.setdefault(doc["_province_key"], {
            "value": doc["_province"],
            "label": doc["_province"],
            "communeCount": 0,
            "provinceCount": 0,
            "missingPasswordCount": 0,
        })
        item["communeCount" if doc.get("role") == "commune" else "provinceCount"] += 1
        if not doc.get("password_enc"):
            item["missingPasswordCount"] += 1
    return {
        "provinces": sorted(provinces.values(), key=lambda item: fold(item["label"])),
        "vaultEnabled": password_vault.enabled(),
    }


def _row(doc: dict) -> list[str]:
    role = doc.get("role")
    return [
        doc.get("username") or "",
        password_vault.decrypt(doc.get("password_enc")) or "",
        doc.get("name") or "",
        doc["_province"],
        # HCC tỉnh là đơn vị cấp tỉnh: dữ liệu cũ có thể còn sót `xa`, không đưa vào danh sách.
        (doc.get("xa") or "") if role == "commune" else "",
        ROLE_LABELS.get(role, role or ""),
        "Tạm khoá" if doc.get("access_disabled") is True else "Hoạt động",
    ]


def _sort_key(doc: dict) -> tuple:
    return (
        doc["_province_key"],
        0 if doc.get("role") == "province" else 1,
        fold(doc.get("xa")),
        doc.get("username") or "",
    )


_SHEET_FORBIDDEN = str.maketrans({ch: " " for ch in '[]:*?/\\'})


def _sheet_title(label: str, used: set[str]) -> str:
    """Excel: tên sheet ≤31 ký tự, không chứa []:*?/\\, không trùng (không phân biệt hoa thường)."""
    base = " ".join(label.translate(_SHEET_FORBIDDEN).split())[:31] or "Tỉnh"
    title, n = base, 2
    while title.casefold() in used:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title.casefold())
    return title


def _fill_sheet(sheet, docs: list[dict]) -> None:
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for doc in docs:
        sheet.append(_row(doc))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            # openpyxl coi chuỗi bắt đầu bằng "=" là CÔNG THỨC; mật khẩu/tên hoàn toàn có thể
            # bắt đầu như vậy. Ép kiểu chữ + định dạng "@" để Excel giữ nguyên, kể cả số 0 đầu.
            cell.data_type = "s"
            cell.number_format = "@"
    for index, width in enumerate(_WIDTHS):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def build_workbook(docs: list[dict]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    groups: dict[str, list[dict]] = {}
    for doc in sorted(docs, key=_sort_key):
        groups.setdefault(doc["_province_key"], []).append(doc)
    used: set[str] = set()
    for group in groups.values():
        _fill_sheet(workbook.create_sheet(_sheet_title(group[0]["_province"], used)), group)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def export(province: str | None) -> tuple[bytes, str, int, str | None]:
    """(nội dung, tên file, số tài khoản, tên tỉnh đã chuẩn hoá hoặc None = tất cả)."""
    docs = await _accounts()
    label = province_name(province) if (province or "").strip() else None
    if label:
        key = fold(label)
        docs = [doc for doc in docs if doc["_province_key"] == key]
        if not docs:
            raise AppError(
                "EXPORT_PROVINCE_EMPTY",
                "Tỉnh/thành đã chọn không có tài khoản Hành chính công nào.",
                400,
            )
    elif not docs:
        raise AppError("EXPORT_EMPTY", "Chưa có tài khoản Hành chính công nào để xuất.", 400)
    today = datetime.now(_VN).strftime("%d-%m-%Y")
    filename = f"Tài khoản HCC - {label or 'Tất cả tỉnh'} - {today}.xlsx"
    return build_workbook(docs), filename, len(docs), label
