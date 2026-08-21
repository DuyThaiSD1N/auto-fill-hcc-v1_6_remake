"""Dựng workbook Excel tổng hợp; không truy cập DB để dễ kiểm thử."""
import io
import re
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.procedures.registry import PROCEDURES
from app.reports.integration import canonical_procedure_id, unit_key
from app.reports.procedure_meta import cap_thu_tuc, ma_thu_tuc, pham_vi_ho_tro
from app.traces.date_range import VIETNAM_TZ


_LOCAL_HEADERS = [
    "STT",
    "Mã thủ tục",
    "Tên thủ tục",
    "Thủ tục thuộc cấp",
    "Phạm vi hỗ trợ",
    "Số lượng hồ sơ đã tiếp nhận",
]
_COMBINED_HEADERS = [
    "STT",
    "Mã thủ tục",
    "Tên thủ tục",
    "Thủ tục thuộc cấp",
    "Phạm vi hỗ trợ",
    "Số lượng hồ sơ đã tiếp nhận (bản chưa tích hợp giọng nói)",
    "Số lượng hồ sơ đã tiếp nhận (bản đã tích hợp giọng nói)",
    "TỔNG SỐ HỒ SƠ",
]
_LOCAL_WIDTHS = [6, 13, 62, 17, 34, 15]
_COMBINED_WIDTHS = [6, 13, 62, 17, 34, 22, 22, 16]
_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _sheet_title(account: dict, used: set[str]) -> str:
    raw = account.get("xa") or account.get("name") or account.get("username") or "Tài khoản"
    base = _INVALID_SHEET_CHARS.sub("-", str(raw)).strip().strip("'") or "Tài khoản"
    candidate = base[:31]
    index = 2
    while candidate.casefold() in used:
        suffix = f" ({index})"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        index += 1
    used.add(candidate.casefold())
    return candidate


def _display_date_range(date_from: datetime, date_to: datetime) -> str:
    first = date_from.astimezone(VIETNAM_TZ).strftime("%d/%m/%Y")
    # date_to là mốc loại trừ đầu ngày kế tiếp.
    last = (date_to - timedelta(microseconds=1)).astimezone(VIETNAM_TZ).strftime("%d/%m/%Y")
    return f"Từ ngày {first} đến hết ngày {last} (giờ Việt Nam)"


def _procedure_rows(
    procedures: list[dict],
    handfree_procedures: list[dict] | None = None,
) -> list[dict]:
    registry = {item["key"]: item for item in PROCEDURES}
    combined = handfree_procedures is not None
    rows_by_id: dict[str, dict] = {}
    for item in procedures:
        count = int(item.get("count") or 0)
        if count <= 0:
            continue
        key = item.get("key") or "—"
        entry = registry.get(key)
        canonical_id = canonical_procedure_id(key, entry)
        row = rows_by_id.setdefault(canonical_id, {
            "ma": ma_thu_tuc(entry),
            "ten": (entry or {}).get("label") or item.get("label") or key,
            "cap": cap_thu_tuc(key),
            "pham_vi": pham_vi_ho_tro(entry, key),
            "auto_count": 0,
            "handfree_count": 0,
        })
        row["auto_count"] += count
    for item in handfree_procedures or []:
        count = int(item.get("count") or 0)
        if count <= 0:
            continue
        key = str(item.get("procedureKey") or "").strip() or "—"
        entry = registry.get(key)
        canonical_id = (
            str(item.get("canonicalProcedureId") or "").strip()
            or canonical_procedure_id(key, entry)
        )
        row = rows_by_id.setdefault(canonical_id, {
            "ma": item.get("procedureCode") or ma_thu_tuc(entry),
            "ten": (entry or {}).get("label") or item.get("label") or key,
            "cap": item.get("level") or cap_thu_tuc(key),
            "pham_vi": item.get("supportScope") or pham_vi_ho_tro(entry, key),
            "auto_count": 0,
            "handfree_count": 0,
        })
        row["handfree_count"] += count

    rows = list(rows_by_id.values())
    for row in rows:
        row["total"] = row["auto_count"] + row["handfree_count"]
        if not combined:
            row["count"] = row["auto_count"]
    rows.sort(key=lambda row: (-row["total"], row["ten"]))
    return rows


def _write_sheet(
    workbook: Workbook,
    *,
    account: dict,
    procedures: list[dict],
    handfree_procedures: list[dict] | None,
    date_from: datetime,
    date_to: datetime,
    used_titles: set[str],
) -> None:
    worksheet = workbook.create_sheet(title=_sheet_title(account, used_titles))
    combined = handfree_procedures is not None
    headers = _COMBINED_HEADERS if combined else _LOCAL_HEADERS
    widths = _COMBINED_WIDTHS if combined else _LOCAL_WIDTHS
    max_column = len(headers)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    unit = account.get("name") or account.get("xa") or account.get("username") or "—"
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_column)
    title = worksheet["A1"]
    title.value = "THỐNG KÊ HỒ SƠ TIẾP NHẬN QUA TRỢ LÝ HỖ TRỢ THỦ TỤC HÀNH CHÍNH"
    title.font = Font(bold=True, size=13)
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 32

    metadata = [
        f"Đơn vị: {unit}",
        f"Địa bàn: {account.get('xa') or '—'} · {account.get('tinh') or 'Chưa xác định tỉnh'}",
        _display_date_range(date_from, date_to),
    ]
    for row_index, value in enumerate(metadata, start=2):
        worksheet.merge_cells(
            start_row=row_index, start_column=1, end_row=row_index, end_column=max_column,
        )
        cell = worksheet.cell(row=row_index, column=1, value=value)
        cell.font = Font(italic=row_index == 4, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    header_row = 7
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[header_row].height = 30
    worksheet.freeze_panes = f"A{header_row + 1}"

    rows = _procedure_rows(procedures, handfree_procedures)
    current_row = header_row
    for index, row in enumerate(rows, start=1):
        current_row = header_row + index
        values = [index, row["ma"], row["ten"], row["cap"], row["pham_vi"]]
        if combined:
            values.extend([row["auto_count"], row["handfree_count"], row["total"]])
        else:
            values.append(row["count"])
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=current_row, column=column, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(
                horizontal="center" if column in (1, 2, 4) or column >= 6 else "left",
                vertical="center",
                wrap_text=True,
            )

    if not rows:
        current_row = header_row + 1
        worksheet.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=max_column,
        )
        cell = worksheet.cell(
            row=current_row,
            column=1,
            value="Không có hồ sơ trong khoảng thời gian đã chọn.",
        )
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_row = current_row + 1
    worksheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    total_label = worksheet.cell(row=total_row, column=1, value="TỔNG CỘNG")
    total_label.font = Font(bold=True)
    total_label.alignment = Alignment(horizontal="center", vertical="center")
    if combined:
        totals = [
            sum(row["auto_count"] for row in rows),
            sum(row["handfree_count"] for row in rows),
            sum(row["total"] for row in rows),
        ]
        for column, value in enumerate(totals, start=6):
            cell = worksheet.cell(row=total_row, column=column, value=value)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    else:
        total_value = worksheet.cell(
            row=total_row, column=6, value=sum(row["count"] for row in rows),
        )
        total_value.font = Font(bold=True)
        total_value.alignment = Alignment(horizontal="center", vertical="center")
    for column in range(1, max_column + 1):
        cell = worksheet.cell(row=total_row, column=column)
        cell.border = _BORDER
        cell.fill = _TOTAL_FILL


def build_excel(
    *,
    accounts: list[dict],
    stats: dict,
    handfree_stats: dict | None = None,
    date_from: datetime,
    date_to: datetime,
) -> bytes:
    """Một tài khoản luôn tạo đúng một sheet, kể cả khi không phát sinh hồ sơ."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    wards = {item["userId"]: item for item in stats.get("wards") or []}
    handfree_units = {
        item.get("unitKey"): item
        for item in (handfree_stats or {}).get("units") or []
        if item.get("unitKey")
    }
    used_titles: set[str] = set()
    for account in accounts:
        account_id = str(account["_id"])
        handfree_unit = handfree_units.get(unit_key(account.get("tinh"), account.get("xa")))
        _write_sheet(
            workbook,
            account=account,
            procedures=(wards.get(account_id) or {}).get("procedures") or [],
            handfree_procedures=(
                (handfree_unit or {}).get("procedures") or []
                if handfree_stats is not None else None
            ),
            date_from=date_from,
            date_to=date_to,
            used_titles=used_titles,
        )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
