"""Nhập tài khoản hàng loạt từ Excel (app/users/import_excel.py).

Chốt nghiệp vụ:
- Username đã có (kể cả đã xóa mềm) → BỎ QUA, không đụng tài khoản cũ.
- Xã chỉ ghi tên trần → phải lưu TÊN ĐẦY ĐỦ tra từ danh mục, so CÓ DẤU trước.
- File chỉ tạo được 2 role nghiệp vụ; không đường nào tạo admin.
- Mật khẩu không bao giờ nằm trong kết quả trả về.
Dữ liệu test là giá trị giả.
"""
import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from pymongo.errors import DuplicateKeyError

from app.core.errors import AppError
from app.users import import_excel

HEADERS = ["STT", "Tên", "Tên tỉnh", "Tên xã/phường", "Tên đăng nhập", "Mật khẩu", "Role"]
XA = "Hành chính công xã"
TINH = "Hành chính công tỉnh"


def _xlsx(rows, headers=HEADERS, title_rows=()):
    wb = Workbook()
    ws = wb.active
    for row in title_rows:
        ws.append(list(row))
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class _Cursor:
    def __init__(self, docs):
        self._docs = list(docs)

    def __aiter__(self):
        self._it = iter(self._docs)
        return self

    async def __anext__(self):
        try:
            return next(self._it)
        except StopIteration:
            raise StopAsyncIteration


class _Users:
    def __init__(self, existing=(), race=()):
        self.docs = [dict(d) for d in existing]
        self.inserted = []
        self._race = set(race)

    def find(self, query, projection=None):
        wanted = set(query["username"]["$in"])
        return _Cursor(d for d in self.docs if d["username"] in wanted)

    async def insert_one(self, doc):
        if doc["username"] in self._race or any(d["username"] == doc["username"] for d in self.docs):
            raise DuplicateKeyError("E11000")
        self.docs.append(doc)
        self.inserted.append(doc)


@pytest.fixture()
def users(monkeypatch):
    holder = {"users": _Users()}
    monkeypatch.setattr(import_excel, "get_db", lambda: type("Db", (), {"users": holder["users"]})())
    # bcrypt thật tốn ~170ms/lần — test chỉ cần biết đã băm, không cần băm thật.
    monkeypatch.setattr(import_excel, "hash_password", lambda p: f"hashed::{len(p)}")

    def use(**kwargs):
        holder["users"] = _Users(**kwargs)
        return holder["users"]

    return use


def _by_user(result):
    return {row["username"]: row for row in result["rows"]}


async def test_tim_tieu_de_duoi_dong_tieu_de_trang_va_bo_qua_cot_la(users):
    users()
    headers = ["STT", "Tên đơn vị", "Mã cơ quan", "Tài khoản đăng nhập Trợ lý nhân dân",
               "Mật khẩu", "Ghi chú", "Tên tỉnh", "Tên xã/phường", "Role"]
    content = _xlsx(
        [[1, "UBND phường A", "H00.1", "tkthu01", "matkhau-gia-01", "", "Bắc Ninh", "Bồng Lai", XA]],
        headers=headers, title_rows=[["DANH SÁCH TÀI KHOẢN"], []],
    )
    row = _by_user(await import_excel.run(content, apply=False))["tkthu01"]
    assert row["status"] == "create", row
    assert row["name"] == "UBND phường A"


async def test_xa_ten_tran_luu_ten_day_du_tu_danh_muc(users):
    users()
    content = _xlsx([[1, "A", "Bắc Ninh", "Bồng Lai", "tkthu01", "matkhau-gia-01", XA]])
    row = (await import_excel.run(content, apply=False))["rows"][0]
    assert (row["tinh"], row["xa"]) == ("Tỉnh Bắc Ninh", "Phường Bồng Lai")


async def test_xa_so_co_dau_truoc_khong_bao_trung_oan(users):
    """Bỏ dấu thì 'Văn Lang' và 'Văn Lăng' (hai xã khác nhau) thành một."""
    users()
    content = _xlsx([
        [1, "A", "Thái Nguyên", "Văn Lăng", "tkthu01", "matkhau-gia-01", XA],
        [2, "B", "Thái Nguyên", "Van Lang", "tkthu02", "matkhau-gia-02", XA],
    ])
    rows = _by_user(await import_excel.run(content, apply=False))
    assert rows["tkthu01"]["xa"] == "Xã Văn Lăng"
    assert rows["tkthu02"]["status"] == "error", "thiếu dấu mà khớp 2 xã thì phải báo, không tự chọn"
    assert "Xã Văn Lang" in rows["tkthu02"]["message"] and "Xã Văn Lăng" in rows["tkthu02"]["message"]


async def test_role_chi_nhan_hai_loai_nghiep_vu(users):
    users()
    content = _xlsx([
        [1, "A", "Bắc Ninh", "", "tkthu01", "matkhau-gia-01", "hanh chinh cong tinh"],
        [2, "B", "Bắc Ninh", "", "tkthu02", "matkhau-gia-02", "admin"],
        [3, "C", "Bắc Ninh", "", "tkthu03", "matkhau-gia-03", "Quản trị viên"],
        [4, "D", "Bắc Ninh", "", "tkthu04", "matkhau-gia-04", "province_admin"],
    ])
    rows = _by_user(await import_excel.run(content, apply=False))
    assert rows["tkthu01"]["role"] == "province" and rows["tkthu01"]["status"] == "create"
    for leo_quyen in ("tkthu02", "tkthu03", "tkthu04"):
        assert rows[leo_quyen]["status"] == "error", "file tải lên không được tạo tài khoản quản trị"


async def test_tai_khoan_xa_bat_buoc_co_xa_tai_khoan_tinh_thi_khong(users):
    users()
    content = _xlsx([
        [1, "Sở X", "Bắc Ninh", "", "tkthu01", "matkhau-gia-01", TINH],
        [2, "UBND", "Bắc Ninh", "", "tkthu02", "matkhau-gia-02", XA],
    ])
    rows = _by_user(await import_excel.run(content, apply=False))
    assert rows["tkthu01"]["status"] == "create" and not rows["tkthu01"]["xa"]
    assert rows["tkthu02"]["status"] == "error"


async def test_loi_tinh_xa_mat_khau_ngan_bao_theo_tung_dong(users):
    users()
    content = _xlsx([
        [1, "A", "Tỉnh Không Có", "", "tkthu01", "matkhau-gia-01", TINH],
        [2, "B", "Bắc Ninh", "Không Có Xã", "tkthu02", "matkhau-gia-02", XA],
        [3, "C", "Bắc Ninh", "Bồng Lai", "tkthu03", "ngan", XA],
    ])
    result = await import_excel.run(content, apply=False)
    assert [r["status"] for r in result["rows"]] == ["error", "error", "error"]
    assert [r["row"] for r in result["rows"]] == [2, 3, 4], "số dòng phải là dòng Excel thật"


async def test_mat_khau_o_so_khong_thanh_so_thuc(users):
    store = users()
    content = _xlsx([[1, "A", "Bắc Ninh", "Bồng Lai", "tkthu01", 12345678, XA]])
    await import_excel.run(content, apply=True)
    assert store.inserted[0]["password_hash"] == "hashed::8", "12345678 không được thành '12345678.0'"


async def test_trung_trong_file_khong_phan_biet_hoa_thuong(users):
    users()
    content = _xlsx([
        [1, "A", "Bắc Ninh", "Bồng Lai", "TkThu01", "matkhau-gia-01", XA],
        [2, "B", "Bắc Ninh", "Bồng Lai", "tkthu01", "matkhau-gia-02", XA],
    ])
    rows = (await import_excel.run(content, apply=False))["rows"]
    assert rows[0]["status"] == "create"
    assert rows[1]["status"] == "duplicate_in_file" and "dòng 2" in rows[1]["message"]


async def test_da_ton_tai_thi_bo_qua_ke_ca_da_xoa_mem(users):
    store = users(existing=[
        {"username": "tkcu01", "deleted_at": None, "password_hash": "cu"},
        {"username": "tkxoa01", "deleted_at": "2026-09-01", "password_hash": "cu"},
    ])
    content = _xlsx([
        [1, "A", "Bắc Ninh", "Bồng Lai", "TKCU01", "matkhau-gia-01", XA],
        [2, "B", "Bắc Ninh", "Bồng Lai", "tkxoa01", "matkhau-gia-02", XA],
        [3, "C", "Bắc Ninh", "Bồng Lai", "tkmoi01", "matkhau-gia-03", XA],
    ])
    result = await import_excel.run(content, apply=True)
    rows = _by_user(result)
    assert rows["tkcu01"]["status"] == "exists"
    assert rows["tkxoa01"]["status"] == "exists_deleted" and "Đã xóa" in rows["tkxoa01"]["message"]
    assert rows["tkmoi01"]["status"] == "created"
    assert [d["username"] for d in store.inserted] == ["tkmoi01"]
    assert store.docs[0]["password_hash"] == "cu", "tài khoản cũ không được bị đổi mật khẩu"


async def test_xem_truoc_khong_ghi_gi(users):
    store = users()
    content = _xlsx([[1, "A", "Bắc Ninh", "Bồng Lai", "tkthu01", "matkhau-gia-01", XA]])
    result = await import_excel.run(content, apply=False)
    assert result["applied"] is False and store.inserted == []
    assert result["summary"] == {"create": 1}


async def test_tao_that_ghi_dung_truong_va_khong_tra_mat_khau(users):
    store = users()
    content = _xlsx([[1, "UBND A", "Bắc Ninh", "Bồng Lai", "tkthu01", "matkhau-gia-01", XA]])
    result = await import_excel.run(content, apply=True)
    doc = store.inserted[0]
    assert {k: doc[k] for k in ("username", "name", "tinh", "xa", "role", "access_disabled")} == {
        "username": "tkthu01", "name": "UBND A", "tinh": "Tỉnh Bắc Ninh",
        "xa": "Phường Bồng Lai", "role": "commune", "access_disabled": False,
    }
    assert "password" not in doc
    assert "matkhau-gia-01" not in json.dumps(result, ensure_ascii=False, default=str)


async def test_trung_do_chay_dua_luc_ghi_thi_bao_da_ton_tai(users):
    users(race={"tkthu01"})
    content = _xlsx([[1, "A", "Bắc Ninh", "Bồng Lai", "tkthu01", "matkhau-gia-01", XA]])
    row = (await import_excel.run(content, apply=True))["rows"][0]
    assert row["status"] == "exists"


async def test_thieu_cot_bat_buoc_va_file_hong(users):
    users()
    no_role = _xlsx([[1, "A", "Bắc Ninh", "", "tkthu01", "matkhau-gia-01"]], headers=HEADERS[:-1])
    with pytest.raises(AppError) as err:
        await import_excel.run(no_role, apply=False)
    assert err.value.error == "IMPORT_MISSING_COLUMNS" and "Role" in err.value.message
    with pytest.raises(AppError) as err:
        await import_excel.run(b"khong phai excel", apply=False)
    assert err.value.error == "IMPORT_BAD_FILE"


def test_file_mau_doc_lai_duoc_va_khong_co_dong_vi_du():
    """Dòng ví dụ trong file mẫu dễ bị nhập nhầm thành tài khoản thật → sheet đầu chỉ có tiêu đề."""
    content = import_excel.template_bytes()
    ws = load_workbook(BytesIO(content)).worksheets[0]
    assert [c.value for c in ws[1]] == HEADERS
    with pytest.raises(AppError) as err:
        import_excel.parse_workbook(content)
    assert err.value.error == "IMPORT_EMPTY", "tiêu đề của file mẫu phải được bộ đọc nhận ra"


# ── Tầng router: multipart thật, file mẫu, giới hạn dung lượng, quyền admin ──
from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.core.deps import require_admin  # noqa: E402
from app.core.errors import app_error_handler  # noqa: E402
from app.users.router import router as users_router  # noqa: E402

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _client(admin=True):
    app = FastAPI()
    app.add_exception_handler(AppError, app_error_handler)
    app.include_router(users_router)
    if admin:
        app.dependency_overrides[require_admin] = lambda: {"id": "admin-gia", "role": "admin"}
    return TestClient(app)


def test_router_xem_truoc_qua_multipart(users):
    users()
    content = _xlsx([[1, "A", "Bắc Ninh", "Bồng Lai", "tkthu01", "matkhau-gia-01", XA]])
    res = _client().post("/api/v1/users/import", files={"file": ("ds.xlsx", content, XLSX_MIME)})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["applied"] is False and body["summary"] == {"create": 1}
    assert "matkhau-gia-01" not in res.text


def test_router_file_mau_tai_ve_la_xlsx():
    res = _client().get("/api/v1/users/import/template")
    assert res.status_code == 200
    assert res.headers["content-type"] == XLSX_MIME
    assert "attachment" in res.headers["content-disposition"]
    assert [c.value for c in load_workbook(BytesIO(res.content)).worksheets[0][1]] == HEADERS


def test_router_chan_file_qua_2mb(users):
    users()
    big = b"x" * (import_excel.MAX_BYTES + 10)
    res = _client().post("/api/v1/users/import", files={"file": ("to.xlsx", big, XLSX_MIME)})
    assert res.status_code == 413


def test_router_khong_phai_admin_thi_khong_vao_duoc(users):
    users()
    content = _xlsx([[1, "A", "Bắc Ninh", "Bồng Lai", "tkthu01", "matkhau-gia-01", XA]])
    client = _client(admin=False)
    assert client.post(
        "/api/v1/users/import?apply=true", files={"file": ("ds.xlsx", content, XLSX_MIME)},
    ).status_code in (401, 403)
    assert client.get("/api/v1/users/import/template").status_code in (401, 403)
