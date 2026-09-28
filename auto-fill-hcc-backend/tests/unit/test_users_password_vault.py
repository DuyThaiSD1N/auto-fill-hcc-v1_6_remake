"""Mật khẩu tài khoản lưu HAI CHIỀU (xem / nạp / xuất) bên cạnh bcrypt.

Chốt:
- Mọi đường đặt mật khẩu ghi cả bcrypt lẫn `password_enc`; không có khoá thì XOÁ `password_enc`
  (giữ bản cũ là hiện một mật khẩu đã sai).
- `password_enc` không bao giờ nằm trong dữ liệu tài khoản trả FE.
- Nạp mật khẩu chỉ lưu dòng KHỚP bcrypt, không đổi mật khẩu của ai.
- Xuất chỉ HCC xã/tỉnh, gom tỉnh theo tên chuẩn hoá, HCC tỉnh luôn trống cột xã, chưa có bản
  lưu thì trống cột mật khẩu.
Dữ liệu test là giá trị giả.
"""
from io import BytesIO

import pytest
from bson import ObjectId
from cryptography.fernet import Fernet
from openpyxl import Workbook, load_workbook

from app.core.errors import AppError
from app.users import export_excel, import_excel, password_import, password_vault, service
from app.users.schemas import UserCreate, UserUpdate


def _match(doc: dict, query: dict) -> bool:
    for key, cond in query.items():
        value = doc.get(key)
        if isinstance(cond, dict):
            if "$in" in cond and value not in cond["$in"]:
                return False
            if "$ne" in cond and value == cond["$ne"]:
                return False
        elif value != cond:
            return False
    return True


class _Cursor:
    def __init__(self, docs):
        self._docs = list(docs)

    def __aiter__(self):
        self._it = iter(self._docs)
        return self

    async def __anext__(self):
        try:
            return next(self._it)
        except StopIteration:
            raise StopAsyncIteration


class _Users:
    def __init__(self, docs=()):
        self.docs = [dict(d) for d in docs]

    def find(self, query, projection=None):
        return _Cursor(dict(d) for d in self.docs if _match(d, query))

    async def find_one(self, query):
        return next((dict(d) for d in self.docs if _match(d, query)), None)

    async def insert_one(self, doc):
        doc = {"_id": ObjectId(), **doc}
        self.docs.append(doc)
        return type("R", (), {"inserted_id": doc["_id"]})()

    def _apply(self, doc, update):
        doc.update(update.get("$set", {}))
        for field in update.get("$unset", {}):
            doc.pop(field, None)

    async def find_one_and_update(self, query, update, return_document=True):
        doc = next((d for d in self.docs if _match(d, query)), None)
        if doc is None:
            return None
        self._apply(doc, update)
        return dict(doc)

    async def update_one(self, query, update):
        doc = next((d for d in self.docs if _match(d, query)), None)
        if doc is not None:
            self._apply(doc, update)
        return type("R", (), {"matched_count": 1 if doc else 0})()


@pytest.fixture()
def key(monkeypatch):
    value = Fernet.generate_key().decode()
    monkeypatch.setattr(password_vault.settings, "password_vault_key", value)
    # bcrypt thật ~170ms/lần — test chỉ cần biết hash đi cùng mật khẩu nào.
    monkeypatch.setattr(password_vault, "hash_password", lambda p: f"hash::{p}")
    return value


@pytest.fixture()
def no_key(monkeypatch):
    monkeypatch.setattr(password_vault.settings, "password_vault_key", "")
    monkeypatch.setattr(password_vault, "hash_password", lambda p: f"hash::{p}")


def _use_db(monkeypatch, module, users):
    db = type("Db", (), {"users": users})()
    monkeypatch.setattr(module, "get_db", lambda: db)
    return users


# ---------- vault ----------

def test_ma_hoa_giai_ma_dung_chieu(key):
    token = password_vault.encrypt("MatKhau#123")
    assert token and "MatKhau" not in token
    assert password_vault.decrypt(token) == "MatKhau#123"


def test_khong_co_khoa_thi_khong_luu_va_xoa_ban_cu(no_key):
    set_fields, unset_fields = password_vault.password_fields("MatKhau#123")
    assert set_fields == {"password_hash": "hash::MatKhau#123"}
    assert unset_fields == {"password_enc": ""}
    assert password_vault.enabled() is False


def test_khoa_sai_dinh_dang_coi_nhu_tat(monkeypatch):
    monkeypatch.setattr(password_vault.settings, "password_vault_key", "khong-phai-khoa")
    assert password_vault.enabled() is False
    assert password_vault.encrypt("abcdefgh") is None


def test_doi_khoa_thi_ban_cu_giai_ma_ra_none(key, monkeypatch):
    token = password_vault.encrypt("abcdefgh")
    monkeypatch.setattr(password_vault.settings, "password_vault_key", Fernet.generate_key().decode())
    assert password_vault.decrypt(token) is None


# ---------- service ----------

def test_public_bao_co_ban_luu_nhung_khong_lo_ban_ma_hoa():
    out = service._public({"_id": ObjectId(), "username": "a", "password_enc": "tok"})
    assert out["password_stored"] is True
    assert "password_enc" not in out and "password_hash" not in out


async def test_tao_tai_khoan_luu_ca_hai(key, monkeypatch):
    users = _use_db(monkeypatch, service, _Users())
    out = await service.create_user(UserCreate(username="hcc_a", password="MatKhau#123", role="user"))
    doc = users.docs[0]
    assert doc["password_hash"] == "hash::MatKhau#123"
    assert password_vault.decrypt(doc["password_enc"]) == "MatKhau#123"
    assert out["password_stored"] is True


async def test_sua_mat_khau_khong_khoa_thi_xoa_ban_cu(no_key, monkeypatch):
    oid = ObjectId()
    users = _use_db(monkeypatch, service, _Users([
        {"_id": oid, "username": "a", "role": "user", "password_hash": "old", "password_enc": "cu"},
    ]))
    await service.update_user(str(oid), UserUpdate(password="MatKhauMoi1"), str(ObjectId()))
    assert users.docs[0]["password_hash"] == "hash::MatKhauMoi1"
    assert "password_enc" not in users.docs[0]


async def test_dat_lai_mat_khau(key, monkeypatch):
    oid = ObjectId()
    users = _use_db(monkeypatch, service, _Users([
        {"_id": oid, "username": "a", "role": "commune", "password_hash": "old"},
    ]))
    out = await service.set_password(str(oid), "MatKhauMoi1")
    assert users.docs[0]["password_hash"] == "hash::MatKhauMoi1"
    assert password_vault.decrypt(users.docs[0]["password_enc"]) == "MatKhauMoi1"
    assert out["password_stored"] is True


async def test_dat_lai_va_xem_mat_khau_chan_tai_khoan_monitor(key, monkeypatch):
    oid = ObjectId()
    _use_db(monkeypatch, service, _Users([
        {"_id": oid, "username": "mon", "role": "super_admin", "password_hash": "x"},
    ]))
    with pytest.raises(AppError) as exc:
        await service.set_password(str(oid), "MatKhauMoi1")
    assert exc.value.error == "PROTECTED_ACCOUNT"
    with pytest.raises(AppError):
        await service.reveal_password(str(oid))


async def test_xem_mat_khau(key, monkeypatch):
    stored, legacy = ObjectId(), ObjectId()
    _use_db(monkeypatch, service, _Users([
        {"_id": stored, "username": "a", "role": "user",
         "password_enc": password_vault.encrypt("MatKhau#123")},
        {"_id": legacy, "username": "b", "role": "user", "password_hash": "bcrypt-only"},
    ]))
    result, _ = await service.reveal_password(str(stored))
    assert result == {"stored": True, "password": "MatKhau#123", "vaultEnabled": True}
    result, _ = await service.reveal_password(str(legacy))
    assert result["stored"] is False and result["password"] is None


async def test_nhap_excel_luu_ban_ma_hoa(key, monkeypatch):
    users = _use_db(monkeypatch, import_excel, _Users())
    monkeypatch.setattr(import_excel, "hash_password", lambda p: f"hash::{p}")
    rows = [{"row": 2, "username": "hcc_x", "name": "", "xa": "", "tinh": "Tỉnh A", "role": "province"}]
    await import_excel._create_rows(rows, {2: "MatKhau#123"})
    assert password_vault.decrypt(users.docs[0]["password_enc"]) == "MatKhau#123"


# ---------- nạp mật khẩu ----------

def _xlsx(rows, headers=("Tên đăng nhập", "Mật khẩu")):
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def pw_users(key, monkeypatch):
    monkeypatch.setattr(password_import, "verify_password", lambda p, h: h == f"hash::{p}")
    return _use_db(monkeypatch, password_import, _Users([
        {"_id": ObjectId(), "username": "khop", "role": "commune", "password_hash": "hash::Dung1234"},
        {"_id": ObjectId(), "username": "sai", "role": "commune", "password_hash": "hash::Khac1234"},
        {"_id": ObjectId(), "username": "roi", "role": "commune", "password_hash": "hash::Roi12345",
         "password_enc": password_vault.encrypt("Roi12345")},
        {"_id": ObjectId(), "username": "daxoa", "role": "commune", "password_hash": "hash::Xoa12345",
         "deleted_at": "x"},
    ]))


async def test_nap_mat_khau_chi_kiem_khong_ghi(pw_users):
    content = _xlsx([
        ("KHOP", "Dung1234"), ("sai", "Dung1234"), ("roi", "Roi12345"),
        ("khongco", "abcdefgh"), ("khop", "Dung1234"), ("daxoa", "Xoa12345"), ("thieu", None),
    ])
    result = await password_import.run(content, apply=False)
    status = {(r["username"], r["row"]): r["status"] for r in result["rows"]}
    assert status[("khop", 2)] == "store"
    assert status[("sai", 3)] == "mismatch"
    assert status[("roi", 4)] == "already"
    assert status[("khongco", 5)] == "not_found"
    assert status[("khop", 6)] == "duplicate_in_file"
    assert status[("daxoa", 7)] == "not_found"
    assert status[("thieu", 8)] == "error"
    assert "Dung1234" not in repr(result)
    assert "password_enc" not in pw_users.docs[0]


async def test_nap_mat_khau_ap_dung_chi_luu_dong_khop(pw_users):
    result = await password_import.run(_xlsx([("khop", "Dung1234"), ("sai", "Dung1234")]), apply=True)
    assert result["summary"] == {"stored": 1, "mismatch": 1}
    assert password_vault.decrypt(pw_users.docs[0]["password_enc"]) == "Dung1234"
    assert pw_users.docs[0]["password_hash"] == "hash::Dung1234"  # không đổi mật khẩu
    assert "password_enc" not in pw_users.docs[1]


async def test_nap_mat_khau_ap_dung_can_khoa(no_key):
    with pytest.raises(AppError) as exc:
        await password_import.run(_xlsx([("a", "abcdefgh")]), apply=True)
    assert exc.value.error == "PASSWORD_VAULT_DISABLED"


# ---------- xuất ----------

@pytest.fixture()
def export_users(key, monkeypatch):
    return _use_db(monkeypatch, export_excel, _Users([
        {"_id": ObjectId(), "username": "dn_xa", "role": "commune", "tinh": "Đà Nẵng",
         "xa": "Phường Hải Châu", "name": "HCC Hải Châu",
         "password_enc": password_vault.encrypt("=MatKhau1")},
        {"_id": ObjectId(), "username": "dn_tinh", "role": "province", "tinh": "Thành phố Đà Nẵng",
         "xa": "Phường Sót Lại", "name": "HCC TP", "access_disabled": True},
        {"_id": ObjectId(), "username": "ld_xa", "role": "commune", "tinh": "Lâm Đồng", "xa": "Xã A"},
        {"_id": ObjectId(), "username": "admin1", "role": "admin", "tinh": "Đà Nẵng"},
        {"_id": ObjectId(), "username": "bc", "role": "province_admin", "tinh": "Đà Nẵng"},
        {"_id": ObjectId(), "username": "daxoa", "role": "commune", "tinh": "Đà Nẵng",
         "xa": "Xã B", "deleted_at": "x"},
        {"_id": ObjectId(), "username": "chuagan", "role": "commune", "tinh": None},
    ]))


def _sheets(data: bytes) -> dict[str, list[list]]:
    wb = load_workbook(BytesIO(data))
    return {
        ws.title: [[c if c is not None else "" for c in row] for row in ws.iter_rows(values_only=True)]
        for ws in wb.worksheets
    }


def _sheet(data: bytes) -> list[list]:
    return next(iter(_sheets(data).values()))


async def test_xuat_options_gom_tinh_va_dem(export_users):
    out = await export_excel.options()
    by_label = {p["label"]: p for p in out["provinces"]}
    assert set(by_label) == {"Thành phố Đà Nẵng", "Tỉnh Lâm Đồng"}
    dn = by_label["Thành phố Đà Nẵng"]
    assert (dn["communeCount"], dn["provinceCount"], dn["missingPasswordCount"]) == (1, 1, 1)


async def test_xuat_theo_tinh(export_users):
    data, filename, count, label = await export_excel.export("Đà Nẵng")
    assert count == 2 and label == "Thành phố Đà Nẵng"
    assert "Thành phố Đà Nẵng" in filename
    rows = _sheet(data)
    assert rows[0] == export_excel.HEADERS
    # HCC tỉnh đứng trước, cột xã trống dù DB còn sót giá trị; chưa lưu mật khẩu → trống.
    assert rows[1] == ["dn_tinh", "", "HCC TP", "Thành phố Đà Nẵng", "", "Hành chính công tỉnh", "Tạm khoá"]
    assert rows[2] == ["dn_xa", "=MatKhau1", "HCC Hải Châu", "Thành phố Đà Nẵng", "Phường Hải Châu",
                       "Hành chính công xã", "Hoạt động"]


async def test_xuat_tat_ca_chi_hcc_moi_tinh_mot_sheet(export_users):
    data, filename, count, label = await export_excel.export(None)
    sheets = _sheets(data)
    assert list(sheets) == ["Thành phố Đà Nẵng", "Tỉnh Lâm Đồng"]
    assert [row[0] for row in sheets["Thành phố Đà Nẵng"][1:]] == ["dn_tinh", "dn_xa"]
    assert [row[0] for row in sheets["Tỉnh Lâm Đồng"][1:]] == ["ld_xa"]
    assert all(rows[0] == export_excel.HEADERS for rows in sheets.values())
    assert count == 3 and label is None and "Tất cả tỉnh" in filename


def test_ten_sheet_hop_le_va_khong_trung():
    used: set[str] = set()
    assert export_excel._sheet_title("A/B:C", used) == "A B C"
    assert export_excel._sheet_title("a b c", used) == "a b c (2)"
    assert len(export_excel._sheet_title("X" * 40, used)) == 31


async def test_xuat_tinh_khong_co_tai_khoan(export_users):
    with pytest.raises(AppError) as exc:
        await export_excel.export("Quảng Ngãi")
    assert exc.value.error == "EXPORT_PROVINCE_EMPTY"


# ---------- router: quyền, không lưu đệm, nhật ký ----------
from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.core.deps import require_admin  # noqa: E402
from app.core.errors import app_error_handler  # noqa: E402
from app.users import access_log  # noqa: E402
from app.users.router import router as users_router  # noqa: E402


@pytest.fixture()
def logs(monkeypatch):
    recorded = []

    async def record(action, actor, **details):
        recorded.append({"action": action, **details})

    monkeypatch.setattr(access_log, "record", record)
    return recorded


def _client(admin=True):
    app = FastAPI()
    app.add_exception_handler(AppError, app_error_handler)
    app.include_router(users_router)
    if admin:
        app.dependency_overrides[require_admin] = lambda: {"id": "admin-gia", "username": "ad", "role": "admin"}
    return TestClient(app)


def test_router_khong_phai_admin_bi_chan():
    client = _client(admin=False)
    oid = str(ObjectId())
    assert client.post(f"/api/v1/users/{oid}/password/reveal").status_code in (401, 403)
    assert client.post(f"/api/v1/users/{oid}/password", json={"password": "abcdefgh"}).status_code in (401, 403)
    assert client.get("/api/v1/users/export/options").status_code in (401, 403)
    assert client.post("/api/v1/users/export", json={}).status_code in (401, 403)
    assert client.get("/api/v1/users/password-import/template").status_code in (401, 403)


def test_router_xem_mat_khau_khong_luu_dem_va_ghi_nhat_ky(key, monkeypatch, logs):
    oid = ObjectId()
    _use_db(monkeypatch, service, _Users([
        {"_id": oid, "username": "a", "role": "commune", "password_enc": password_vault.encrypt("MatKhau#123")},
    ]))
    res = _client().post(f"/api/v1/users/{oid}/password/reveal")
    assert res.status_code == 200 and res.json()["password"] == "MatKhau#123"
    assert res.headers["cache-control"] == "no-store"
    assert logs == [{"action": "reveal_password", "target_id": str(oid), "target_username": "a", "stored": True}]
    assert "MatKhau" not in repr(logs)


def test_router_doi_mat_khau_ngan_bi_tu_choi(key, monkeypatch, logs):
    res = _client().post(f"/api/v1/users/{ObjectId()}/password", json={"password": "ngan"})
    assert res.status_code == 422 and logs == []


def test_router_xuat_excel(export_users, logs):
    res = _client().post("/api/v1/users/export", json={"province": "Lâm Đồng"})
    assert res.status_code == 200
    assert res.headers["cache-control"] == "no-store"
    assert "attachment" in res.headers["content-disposition"]
    assert [row[0] for row in _sheet(res.content)[1:]] == ["ld_xa"]
    assert logs == [{"action": "export_accounts", "province": "Tỉnh Lâm Đồng", "count": 1}]
