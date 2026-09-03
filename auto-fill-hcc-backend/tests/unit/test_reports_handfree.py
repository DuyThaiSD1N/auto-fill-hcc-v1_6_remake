import io

import pytest
from openpyxl import load_workbook

from app.reports.excel import build_daily_excel, build_excel
from app.reports import handfree_client
from app.reports.integration import unit_key
from app.reports.schemas import ExcelExportRequest
from app.traces.date_range import parse_stats_range


def test_combined_excel_unions_procedures_zero_fills_and_totals():
    date_from, date_to = parse_stats_range("2026-08-18", "2026-08-18")
    account = {
        "_id": "auto-1",
        "username": "hccnghiahung",
        "name": "HCC Xã Nghĩa Hưng",
        "xa": "Xã Nghĩa Hưng",
        "tinh": "Tỉnh Ninh Bình",
        "role": "commune",
    }
    content = build_excel(
        accounts=[account],
        stats={"wards": [{
            "userId": "auto-1",
            "procedures": [
                {"key": "chung-thuc-ban-sao", "count": 7},
                {"key": "khai-tu", "count": 1},
            ],
        }]},
        handfree_stats={"source": "handfree", "units": [{
            "unitKey": unit_key("Tỉnh Ninh Bình", "Xã Nghĩa Hưng"),
            "province": "Tỉnh Ninh Bình",
            "ward": "Xã Nghĩa Hưng",
            "procedures": [
                {
                    "canonicalProcedureId": "tthc:2.000815",
                    "procedureKey": "chung-thuc-ban-sao",
                    "procedureCode": "2.000815",
                    "label": "Chứng thực bản sao",
                    "count": 31,
                },
                {
                    "canonicalProcedureId": "tthc:1.000894",
                    "procedureKey": "ket-hon",
                    "procedureCode": "1.000894",
                    "label": "Thủ tục đăng ký kết hôn",
                    "count": 1,
                },
            ],
        }]},
        date_from=date_from,
        date_to=date_to,
    )
    sheet = load_workbook(io.BytesIO(content)).active

    assert sheet["F7"].value == "Số lượng hồ sơ đã tiếp nhận (bản chưa tích hợp giọng nói)"
    assert sheet["G7"].value == "Số lượng hồ sơ đã tiếp nhận (bản đã tích hợp giọng nói)"
    assert sheet["H7"].value == "TỔNG SỐ HỒ SƠ"
    rows = {sheet.cell(row=row, column=2).value: row for row in range(8, 11)}
    assert [sheet.cell(rows["2.000815"], column=column).value for column in (6, 7, 8)] == [7, 31, 38]
    assert [sheet.cell(rows["1.000656"], column=column).value for column in (6, 7, 8)] == [1, 0, 1]
    assert [sheet.cell(rows["1.000894"], column=column).value for column in (6, 7, 8)] == [0, 1, 1]
    assert [sheet.cell(row=11, column=column).value for column in (6, 7, 8)] == [8, 32, 40]


def test_combined_excel_keeps_auto_account_without_handfree_unit():
    date_from, date_to = parse_stats_range("2026-08-18", "2026-08-18")
    account = {
        "_id": "auto-province",
        "username": "hccninhbinh",
        "name": "HCC Tỉnh Ninh Bình",
        "xa": None,
        "tinh": "Tỉnh Ninh Bình",
        "role": "province",
    }
    content = build_excel(
        accounts=[account],
        stats={"wards": [{
            "userId": "auto-province",
            "procedures": [{"key": "khai-tu", "count": 2}],
        }]},
        handfree_stats={"source": "handfree", "units": []},
        date_from=date_from,
        date_to=date_to,
    )
    sheet = load_workbook(io.BytesIO(content)).active

    assert [sheet.cell(row=8, column=column).value for column in (6, 7, 8)] == [2, 0, 2]
    assert [sheet.cell(row=9, column=column).value for column in (6, 7, 8)] == [2, 0, 2]


@pytest.mark.asyncio
async def test_handfree_report_reads_shared_db_with_exact_experience(monkeypatch):
    captured = {}

    async def fake_stats(*, user_ids, date_from, date_to, experience):
        captured.update({
            "user_ids": user_ids,
            "date_from": date_from,
            "date_to": date_to,
            "experience": experience,
        })
        return {"wards": [{
            "userId": "auto-1",
            "procedures": [{"key": "ket-hon", "count": 3}],
        }]}

    monkeypatch.setattr(handfree_client.traces_repo, "stats_by_user_ids", fake_stats)
    body = ExcelExportRequest(
        dateFrom="2026-08-18",
        dateTo="2026-08-18",
        selectionMode="accounts",
        accountIds=["auto-1"],
        includeHandfree=True,
    )

    result = await handfree_client.fetch_handfree_stats([
        {
            "_id": "auto-1",
            "username": "HCCNghiaHung",
            "xa": "Xã Nghĩa Hưng",
            "tinh": "Ninh Bình",
        },
        {
            "_id": "auto-province",
            "username": "HCCNinhBinh",
            "xa": None,
            "tinh": "Ninh Bình",
        },
    ], body)

    assert captured["user_ids"] == ["auto-1"]
    assert captured["experience"] == "handfree"
    assert captured["date_from"].isoformat() == "2026-08-17T17:00:00+00:00"
    assert captured["date_to"].isoformat() == "2026-08-18T17:00:00+00:00"
    assert result["source"] == "handfree"
    assert result["units"][0]["unitKey"] == "tinh ninh binh::xa nghia hung"
    assert result["units"][0]["matchedAccountCount"] == 1
    assert result["units"][0]["procedures"][0]["count"] == 3


@pytest.mark.asyncio
async def test_handfree_report_does_not_query_when_no_joinable_unit(monkeypatch):
    async def fail_if_called(*args, **kwargs):
        raise AssertionError("Không được gọi Handfree khi không có đơn vị đủ khóa ghép")

    monkeypatch.setattr(handfree_client.traces_repo, "stats_by_user_ids", fail_if_called)
    body = ExcelExportRequest(
        dateFrom="2026-08-18",
        dateTo="2026-08-18",
        selectionMode="accounts",
        accountIds=["auto-province"],
        includeHandfree=True,
    )

    result = await handfree_client.fetch_handfree_stats([{
        "_id": "auto-province",
        "username": "HCCNinhBinh",
        "xa": None,
        "tinh": "Ninh Bình",
    }], body)

    assert result == {"source": "handfree", "units": []}


@pytest.mark.asyncio
async def test_handfree_daily_report_uses_one_local_aggregation(monkeypatch):
    calls: list[dict] = []

    async def fake_daily(**kwargs):
        calls.append(kwargs)
        return [
            {"userId": "auto-1", "date": "2026-08-18", "count": 5},
            {"userId": "auto-1", "date": "2026-08-19", "count": 7},
        ]

    monkeypatch.setattr(
        handfree_client.traces_repo,
        "daily_dossier_counts_by_user_ids",
        fake_daily,
    )
    body = ExcelExportRequest(
        dateFrom="2026-08-18",
        dateTo="2026-08-19",
        selectionMode="accounts",
        accountIds=["auto-1"],
        reportLayout="daily_summary",
        includeHandfree=True,
    )

    result = await handfree_client.fetch_handfree_daily_stats([{
        "_id": "auto-1",
        "username": "hccpbacgiang",
        "xa": "Phường Bắc Giang",
        "tinh": "Tỉnh Bắc Ninh",
    }], body)

    assert len(calls) == 1
    assert calls[0]["user_ids"] == ["auto-1"]
    assert calls[0]["experience"] == "handfree"
    assert result["units"][0]["dailyCounts"] == [
        {"date": "2026-08-18", "count": 5},
        {"date": "2026-08-19", "count": 7},
    ]


def test_daily_excel_sums_auto_fill_and_handfree_by_unit_and_day():
    date_from, date_to = parse_stats_range("2026-08-18", "2026-08-19")
    key = unit_key("Tỉnh Bắc Ninh", "Phường Bắc Giang")
    content = build_daily_excel(
        accounts=[{
            "_id": "auto-1",
            "username": "hccpbacgiang",
            "name": "Phường Bắc Giang",
            "xa": "Phường Bắc Giang",
            "tinh": "Tỉnh Bắc Ninh",
            "role": "commune",
        }],
        daily_counts=[
            {"userId": "auto-1", "date": "2026-08-18", "count": 58},
            {"userId": "auto-1", "date": "2026-08-19", "count": 51},
        ],
        handfree_daily_stats={
            "source": "handfree",
            "units": [{
                "unitKey": key,
                "dailyCounts": [
                    {"date": "2026-08-18", "count": 4},
                    {"date": "2026-08-19", "count": 6},
                ],
            }],
        },
        date_from=date_from,
        date_to=date_to,
    )
    sheet = load_workbook(io.BytesIO(content)).active

    assert [sheet.cell(row=5, column=column).value for column in range(1, 6)] == [
        1, "Phường Bắc Giang", 62, 57, 119,
    ]
