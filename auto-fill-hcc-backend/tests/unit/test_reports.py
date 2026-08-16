import io
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.core.errors import AppError
from app.reports import service as report_service
from app.reports.excel import build_excel
from app.reports.schemas import ExcelExportRequest
from app.reports.service import build_options, select_accounts
from app.traces import repo as traces_repo
from app.traces.date_range import parse_stats_range


def _account(
    account_id: str,
    username: str,
    *,
    xa: str,
    tinh: str,
    name: str | None = None,
    role: str = "commune",
) -> dict:
    return {
        "_id": account_id,
        "username": username,
        "name": name,
        "xa": xa,
        "tinh": tinh,
        "role": role,
    }


def test_report_options_only_lists_provinces_present_in_accounts_and_merges_legacy_names():
    result = build_options([
        _account("u1", "tanphong", xa="Tân Phong", tinh="Lai Châu"),
        _account("u2", "taleng", xa="Tả Lèng", tinh="Tỉnh Lai Châu"),
        _account("u3", "legacy", xa="Đơn vị cũ", tinh="Tỉnh chưa có trong danh mục"),
        {"_id": "u4", "username": "no-place", "role": "user"},
        _account("u5", "tester", xa="Tân Phong", tinh="Lai Châu", role="user"),
    ])

    provinces = {item["value"]: item for item in result["provinces"]}
    assert provinces == {
        "Tỉnh Lai Châu": {
            "value": "Tỉnh Lai Châu",
            "label": "Tỉnh Lai Châu",
            "accountCount": 3,
            "officialAccountCount": 2,
        },
        "Tỉnh chưa có trong danh mục": {
            "value": "Tỉnh chưa có trong danh mục",
            "label": "Tỉnh chưa có trong danh mục",
            "accountCount": 1,
            "officialAccountCount": 1,
        },
    }
    assert len(result["accounts"]) == 5
    assert result["accounts"][0]["tinh"] == "Tỉnh Lai Châu"


def test_select_accounts_by_province_or_exact_account_ids():
    accounts = [
        _account("u1", "tanphong", xa="Tân Phong", tinh="Lai Châu"),
        _account("u2", "taleng", xa="Tả Lèng", tinh="Tỉnh Lai Châu"),
        _account("u3", "songlieu", xa="Song Liễu", tinh="Bắc Ninh"),
    ]
    province_body = ExcelExportRequest(
        dateFrom="2026-08-11",
        dateTo="2026-08-11",
        selectionMode="province",
        province="Tỉnh Lai Châu",
    )
    account_body = ExcelExportRequest(
        dateFrom="2026-08-11",
        dateTo="2026-08-11",
        selectionMode="accounts",
        accountIds=["u3", "u1", "u3"],
    )

    assert [item["_id"] for item in select_accounts(accounts, province_body)] == ["u1", "u2"]
    assert [item["_id"] for item in select_accounts(accounts, account_body)] == ["u3", "u1"]
    # Cờ này chỉ có nghĩa ở chế độ theo tỉnh; chọn đích danh tài khoản không bị đổi phạm vi.
    assert [
        item["_id"]
        for item in select_accounts(accounts, account_body.model_copy(update={"officialOnly": True}))
    ] == ["u3", "u1"]

    missing_body = account_body.model_copy(update={"accountIds": ["deleted-user"]})
    with pytest.raises(AppError) as exc:
        select_accounts(accounts, missing_body)
    assert exc.value.error == "REPORT_ACCOUNT_NOT_FOUND"


def test_official_only_province_keeps_hcc_roles_and_excludes_admin_user_and_legacy():
    accounts = [
        _account("u-commune", "commune", xa="Tân Phong", tinh="Lai Châu", role="commune"),
        _account("u-province", "province", xa="Tân Phong", tinh="Lai Châu", role="province"),
        _account("u-user", "tester", xa="Tân Phong", tinh="Lai Châu", role="user"),
        _account("u-admin", "admin", xa="Tân Phong", tinh="Lai Châu", role="admin"),
        {"_id": "u-legacy", "username": "legacy", "xa": "Tân Phong", "tinh": "Lai Châu"},
    ]
    all_body = ExcelExportRequest(
        dateFrom="2026-08-11",
        dateTo="2026-08-11",
        selectionMode="province",
        province="Tỉnh Lai Châu",
    )
    official_body = all_body.model_copy(update={"officialOnly": True})

    assert [item["_id"] for item in select_accounts(accounts, all_body)] == [
        "u-commune",
        "u-province",
        "u-user",
        "u-admin",
        "u-legacy",
    ]
    assert [item["_id"] for item in select_accounts(accounts, official_body)] == [
        "u-commune",
        "u-province",
    ]


def test_official_only_province_reports_clear_error_when_no_hcc_account():
    accounts = [
        _account("u-user", "tester", xa="Song Liễu", tinh="Bắc Ninh", role="user"),
        _account("u-admin", "admin", xa="Song Liễu", tinh="Bắc Ninh", role="admin"),
    ]
    body = ExcelExportRequest(
        dateFrom="2026-08-11",
        dateTo="2026-08-11",
        selectionMode="province",
        province="Tỉnh Bắc Ninh",
        officialOnly=True,
    )

    with pytest.raises(AppError) as exc:
        select_accounts(accounts, body)
    assert exc.value.error == "REPORT_PROVINCE_OFFICIAL_EMPTY"
    assert "hành chính công" in exc.value.message


def test_excel_has_one_safe_unique_sheet_per_account_including_empty_account():
    date_from, date_to = parse_stats_range("2026-08-11", "2026-08-11")
    accounts = [
        _account("u1", "songlieu-a", xa="Song/Liễu", tinh="Bắc Ninh", name="HCC Song Liễu A"),
        _account("u2", "songlieu-b", xa="Song?Liễu", tinh="Tỉnh Bắc Ninh", name="HCC Song Liễu B"),
    ]
    stats = {
        "wards": [
            {
                "userId": "u1",
                "procedures": [{"key": "chung-thuc-ban-sao", "label": "Tên cũ", "count": 3}],
            }
        ]
    }

    content = build_excel(
        accounts=accounts,
        stats=stats,
        date_from=date_from,
        date_to=date_to,
    )
    workbook = load_workbook(io.BytesIO(content))

    assert workbook.sheetnames == ["Song-Liễu", "Song-Liễu (2)"]
    first = workbook[workbook.sheetnames[0]]
    second = workbook[workbook.sheetnames[1]]
    assert first["A2"].value == "Đơn vị: HCC Song Liễu A"
    assert first["A3"].value == "Địa bàn: Song/Liễu · Bắc Ninh"
    assert first["A4"].value == "Từ ngày 11/08/2026 đến hết ngày 11/08/2026 (giờ Việt Nam)"
    assert all("Tài khoản:" not in str(first.cell(row=row, column=1).value) for row in range(1, 8))
    assert first["A7"].value == "STT"
    assert first["F8"].value == 3
    assert first["F9"].value == 3
    assert second["A8"].value == "Không có hồ sơ trong khoảng thời gian đã chọn."
    assert second["F9"].value == 0


@pytest.mark.asyncio
async def test_report_stats_query_uses_exact_user_ids_and_half_open_date_range(monkeypatch):
    captured: dict = {}

    class AggregateResult:
        async def to_list(self, *, length: int):
            captured["length"] = length
            return []

    class Traces:
        def aggregate(self, pipeline, *, allowDiskUse: bool):
            captured["pipeline"] = pipeline
            captured["allowDiskUse"] = allowDiskUse
            return AggregateResult()

    class Database:
        traces = Traces()

    monkeypatch.setattr(traces_repo, "get_db", lambda: Database())
    start = datetime(2026, 8, 10, 17, tzinfo=timezone.utc)
    end = datetime(2026, 8, 11, 17, tzinfo=timezone.utc)

    result = await traces_repo.stats_by_user_ids(
        user_ids=["u1", "u2", "u1"],
        date_from=start,
        date_to=end,
    )

    assert captured["pipeline"][0] == {
        "$match": {
            "user_id": {"$in": ["u1", "u2"]},
            "created_at": {"$gte": start, "$lt": end},
        }
    }
    assert captured["allowDiskUse"] is True
    assert result["totalDossiers"] == 0


@pytest.mark.asyncio
async def test_export_service_selects_accounts_queries_stats_and_returns_named_workbook(monkeypatch):
    accounts = [
        _account("u1", "tanphong", xa="Tân Phong", tinh="Lai Châu"),
        _account("u2", "taleng", xa="Tả Lèng", tinh="Tỉnh Lai Châu"),
    ]
    captured: dict = {}

    async def fake_accounts():
        return accounts

    async def fake_stats(*, user_ids, date_from, date_to):
        captured.update({"user_ids": user_ids, "date_from": date_from, "date_to": date_to})
        return {"wards": []}

    monkeypatch.setattr(report_service, "_all_accounts", fake_accounts)
    monkeypatch.setattr(report_service.traces_repo, "stats_by_user_ids", fake_stats)
    body = ExcelExportRequest(
        dateFrom="2026-08-11",
        dateTo="2026-08-11",
        selectionMode="province",
        province="Tỉnh Lai Châu",
        officialOnly=True,
    )

    content, filename = await report_service.export_excel(body)
    workbook = load_workbook(io.BytesIO(content))

    assert captured["user_ids"] == ["u1", "u2"]
    assert captured["date_from"] == datetime(2026, 8, 10, 17, tzinfo=timezone.utc)
    assert captured["date_to"] == datetime(2026, 8, 11, 17, tzinfo=timezone.utc)
    assert workbook.sheetnames == ["Tân Phong", "Tả Lèng"]
    assert filename == "bao_cao_ho_so_tinh_lai_chau_hcc_2026-08-11_2026-08-11.xlsx"
