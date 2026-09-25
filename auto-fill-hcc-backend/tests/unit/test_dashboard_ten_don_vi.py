"""Dashboard hiển thị đơn vị theo TÊN TÀI KHOẢN (`name`), không theo tên phường/xã (`xa`).

Một phường có thể có nhiều tài khoản: lấy `xa` làm nhãn thì các dòng trùng tên, không phân biệt
được. `xa` chỉ là dự phòng khi tài khoản chưa đặt tên.
"""
from openpyxl import Workbook

from app.dashboard import export


def _ws_units(units):
    wb = Workbook()
    export._sheet_units(wb, {"units": units})
    ws = wb[wb.sheetnames[-1]]
    return [ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)]


def test_bang_don_vi_lay_ten_tai_khoan_truoc():
    labels = _ws_units([
        {"name": "UBND phường Hiệp Hòa", "xa": "Phường Hiệp Hòa", "dossiers": 5},
        {"name": "Bộ phận một cửa Hiệp Hòa", "xa": "Phường Hiệp Hòa", "dossiers": 3},
        {"name": "", "xa": "Xã Kép", "dossiers": 1},
    ])
    assert "UBND phường Hiệp Hòa" in labels and "Bộ phận một cửa Hiệp Hòa" in labels
    assert "Xã Kép" in labels, "chưa đặt tên thì lùi về tên xã"
    assert labels.count("Phường Hiệp Hòa") == 0


def test_ten_file_theo_ten_tai_khoan():
    name = export.export_filename({
        "scope": {}, "range": {},
        "selected": {"name": "UBND phường Hiệp Hòa", "xa": "Phường Hiệp Hòa"},
    })
    assert "UBND phường Hiệp Hòa" in name
