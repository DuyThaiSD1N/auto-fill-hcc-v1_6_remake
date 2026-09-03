import io
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.core.errors import AppError
from app.reports import service as report_service
from app.reports.excel import build_excel
from app.reports.schemas import ExcelExportRequest
from app.reports.service import build_options, select_accounts
from app.traces import repo as traces_repo
from app.traces.date_range import parse_stats_range


def _account(
    account_id: str,
    username: str,
    *,
    xa: str,
    tinh: str,
    name: str | None = None,
    role: str = "commune",
) -> dict:
    return {
        "_id": account_id,
        "username": username,
        "name": name,
        "xa": xa,
        "tinh": tinh,
        "role": role,
    }


def test_stats_range_uses_complete_vietnam_days():
    start, end = parse_stats_range("2026-08-11", "2026-08-11")
    assert start == datetime(2026, 8, 10, 17, tzinfo=timezone.utc)
    assert end == datetime(2026, 8, 11, 17, tzinfo=timezone.utc)


def test_report_options_merge_short_and_official_province_names():
    result = build_options([
        _account("u1", "tanphong", xa="Tân Phong", tinh="Lai Châu"),
        _account("u2", "taleng", xa="Tả Lèng", tinh="Tỉnh Lai Châu"),
        _account("u3", "tester", xa="Tân Phong", tinh="Lai Châu", role="user"),
    ])
    assert result["provinces"] == [{
        "value": "Tỉnh Lai Châu",
        "label": "Tỉnh Lai Châu",
        "accountCount": 3,
        "officialAccountCount": 2,
    }]


def test_official_province_selection_excludes_admin_and_regular_users():
    accounts = [
        _account("u1", "commune", xa="Tân Phong", tinh="Lai Châu", role="commune"),
        _account("u2", "province", xa="Tân Phong", tinh="Lai Châu", role="province"),
        _account("u3", "tester", xa="Tân Phong", tinh="Lai Châu", role="user"),
        _account("u4", "admin", xa="Tân Phong", tinh="Lai Châu", role="admin"),
    ]
    body = ExcelExportRequest(
        dateFrom="2026-08-11",
        dateTo="2026-08-11",
        selectionMode="province",
        province="Tỉnh Lai Châu",
        officialOnly=True,
    )
    assert [item["_id"] for item in select_accounts(accounts, body)] == ["u1", "u2"]


def test_missing_named_account_is_rejected():
    body = ExcelExportRequest(
        dateFrom="2026-08-11",
        dateTo="2026-08-11",
        selectionMode="accounts",
        accountIds=["deleted"],
    )
    with pytest.raises(AppError) as exc:
        select_accounts([], body)
    assert exc.value.error == "REPORT_ACCOUNT_NOT_FOUND"


def test_excel_has_safe_unique_sheet_for_every_selected_account():
    start, end = parse_stats_range("2026-08-11", "2026-08-11")
    accounts = [
        _account("u1", "a", xa="Song/Liễu", tinh="Bắc Ninh", name="HCC A"),
        _account("u2", "b", xa="Song?Liễu", tinh="Bắc Ninh", name="HCC B"),
    ]
    content = build_excel(
        accounts=accounts,
        stats={"wards": [{
            "userId": "u1",
            "procedures": [{"key": "chung-thuc-ban-sao", "label": "Tên cũ", "count": 3}],
        }]},
        date_from=start,
        date_to=end,
    )
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Song-Liễu", "Song-Liễu (2)"]
    assert workbook["Song-Liễu"]["F8"].value == 3
    assert workbook["Song-Liễu (2)"]["A8"].value == "Không có hồ sơ trong khoảng thời gian đã chọn."


@pytest.mark.asyncio
async def test_report_queries_exact_user_ids_and_half_open_range(monkeypatch):
    captured: dict = {}

    async def fake_accounts():
        return [_account("u1", "tanphong", xa="Tân Phong", tinh="Lai Châu")]

    async def fake_stats(*, user_ids, date_from, date_to):
        captured.update({"user_ids": user_ids, "date_from": date_from, "date_to": date_to})
        return {"wards": []}

    monkeypatch.setattr(report_service, "_all_accounts", fake_accounts)
    monkeypatch.setattr(report_service.traces_repo, "stats_by_user_ids", fake_stats)
    body = ExcelExportRequest(
        dateFrom="2026-08-11",
        dateTo="2026-08-11",
        selectionMode="province",
        province="Tỉnh Lai Châu",
        officialOnly=True,
    )
    content, filename = await report_service.export_excel(body)
    assert load_workbook(io.BytesIO(content)).sheetnames == ["Tân Phong"]
    assert captured == {
        "user_ids": ["u1"],
        "date_from": datetime(2026, 8, 10, 17, tzinfo=timezone.utc),
        "date_to": datetime(2026, 8, 11, 17, tzinfo=timezone.utc),
    }
    assert filename == "bao_cao_ho_so_tinh_lai_chau_hcc_2026-08-11_2026-08-11.xlsx"


@pytest.mark.asyncio
async def test_stats_by_user_ids_builds_strict_query(monkeypatch):
    captured: dict = {}

    class TraceCursor:
        def __aiter__(self):
            return self

        async def __anext__(self):
            raise StopAsyncIteration

    class Traces:
        def find(self, query, projection):
            captured["query"] = query
            captured["projection"] = projection
            return TraceCursor()

    class Database:
        traces = Traces()

    monkeypatch.setattr(traces_repo, "get_db", lambda: Database())
    start = datetime(2026, 8, 10, 17, tzinfo=timezone.utc)
    end = datetime(2026, 8, 11, 17, tzinfo=timezone.utc)
    result = await traces_repo.stats_by_user_ids(
        user_ids=["u1", "u2", "u1"],
        date_from=start,
        date_to=end,
    )
    assert captured["query"] == {
        "user_id": {"$in": ["u1", "u2"]},
        "created_at": {"$gte": start, "$lt": end},
    }
    assert result["totalDossiers"] == 0
