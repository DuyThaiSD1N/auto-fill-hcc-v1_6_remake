"""Giấy ủy quyền soạn tại quầy: đọc căn cước hai bên → kiểm tra → xuất .docx / .pdf.

Không phải thủ tục DVC: không registry, không hồ sơ, không lưu dữ liệu đã điền.
"""
import io
import json
from unittest.mock import AsyncMock

import docx
import fitz
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.channels.handfree.authorization_letter import extract as ex
from app.channels.handfree.authorization_letter import router as rt
from app.channels.handfree.chat import flow
from app.core.deps import require_auth
from app.core.errors import AppError, app_error_handler

USER = {"id": "u1", "username": "hccvuninh", "tinh": "Tỉnh Bắc Ninh", "xa": "Phường Vũ Ninh"}


# ── Chuẩn hoá: không đoán nội dung, chỉ định dạng + tô vàng ô cần xem lại ──

def _person(**kw):
    base = {"hoTen": "NGUYỄN VĂN A", "ngaySinh": "1/2/1960", "soDinhDanh": "001060000001",
            "ngayCap": "10-08-2021", "noiCap": "CỤC TRƯỞNG CỤC CẢNH SÁT QUẢN LÝ HÀNH CHÍNH VỀ TRẬT TỰ XÃ HỘI",
            "noiThuongTru": "Thôn Đông, xã X, huyện Y, tỉnh Z", "loaiGiay": "can_cuoc",
            "tepNguon": [0, 2], "chuaChac": []}
    base.update(kw)
    return base


def test_chuan_hoa_ngay_va_noi_cap():
    p = ex.normalize_person(_person(), 4)
    assert p["ngaySinh"] == "01/02/1960" and p["ngayCap"] == "10/08/2021"
    assert p["noiCap"] == "Cục Cảnh sát quản lý hành chính về trật tự xã hội"
    assert p["tepNguon"] == [0, 2] and p["chuaChac"] == []


def test_so_can_cuoc_sai_do_dai_thi_to_vang():
    assert "soDinhDanh" in ex.normalize_person(_person(soDinhDanh="00106000001"), 4)["chuaChac"]


def test_thieu_mat_sau_de_trong_va_to_vang_khong_bia():
    p = ex.normalize_person(_person(ngayCap="", noiCap=""), 4)
    assert p["ngayCap"] == "" and p["noiCap"] == "", "không có ngày cấp thì không suy được nơi cấp"
    assert {"ngayCap", "noiCap"} <= set(p["chuaChac"])


def test_noi_cap_can_cuoc_theo_luat_tu_ngay_cap_van_to_vang():
    p = ex.normalize_person(_person(ngayCap="05/09/2024", noiCap=""), 4)
    assert p["noiCap"] == "Bộ Công an" and "noiCap" in p["chuaChac"]


def test_cmnd_khong_tu_suy_noi_cap():
    p = ex.normalize_person(_person(loaiGiay="cmnd", soDinhDanh="123456789", noiCap=""), 4)
    assert p["noiCap"] == "" and "soDinhDanh" not in p["chuaChac"]


def test_llm_bao_chua_chac_giu_nguyen():
    assert ex.normalize_person(_person(chuaChac=["noiThuongTru", "la"]), 4)["chuaChac"] == ["noiThuongTru"]


def test_bo_nguoi_rong_va_liet_ke_tep_khong_dung():
    out = ex.normalize_result({"people": [_person(tepNguon=[0, 1]), {"hoTen": "", "tepNguon": [3]}]}, 4)
    assert len(out["people"]) == 1
    assert out["unusedFiles"] == [2, 3]


async def test_doc_hai_nguoi_tu_ocr_va_llm(monkeypatch):
    monkeypatch.setattr(ex.ocr, "ocr_per_file", AsyncMock(return_value=[
        {"text": "CĂN CƯỚC CÔNG DÂN 001060000001 NGUYỄN VĂN A"}, {"text": "IDVNM0600000011..."},
        {"text": "CĂN CƯỚC 001190000002 TRẦN THỊ B"}, {"text": "", "error": "OCR lỗi"},
    ]))
    llm = AsyncMock(return_value=json.dumps({"people": [
        _person(tepNguon=[0, 1]),
        _person(hoTen="TRẦN THỊ B", soDinhDanh="001190000002", tepNguon=[2], ngayCap="", noiCap=""),
    ]}))
    monkeypatch.setattr(ex.client, "chat", llm)
    out = await ex.extract_people([{"name": f"{i}.jpg", "type": "image/jpeg", "dataUrl": "x"} for i in range(4)])
    assert [p["hoTen"] for p in out["people"]] == ["NGUYỄN VĂN A", "TRẦN THỊ B"]
    assert out["ocrFailed"] == [3] and out["unusedFiles"] == [3]
    user_msg = llm.await_args.args[0][1]["content"]
    assert "### Tệp 3\n(OCR không đọc được tệp này)" in user_msg


async def test_llm_hong_json_khong_vo(monkeypatch):
    monkeypatch.setattr(ex.ocr, "ocr_per_file", AsyncMock(return_value=[{"text": "CCCD"}]))
    llm = AsyncMock(return_value="xin lỗi")
    monkeypatch.setattr(ex.client, "chat", llm)
    out = await ex.extract_people([{"name": "a", "type": "image/jpeg", "dataUrl": "x"}])
    assert out["people"] == [] and out["unusedFiles"] == [0]
    assert llm.await_count == 2, "thử lại đúng một lần rồi mới bỏ"


async def test_llm_hong_lan_dau_thu_lai_duoc(monkeypatch):
    monkeypatch.setattr(ex.ocr, "ocr_per_file", AsyncMock(return_value=[{"text": "CCCD"}, {"text": "mặt sau"}]))
    monkeypatch.setattr(ex.client, "chat", AsyncMock(side_effect=["…", json.dumps({"people": [_person(tepNguon=[0, 1])]})]))
    out = await ex.extract_people([{"name": "a", "type": "image/jpeg", "dataUrl": "x"}] * 2)
    assert [p["hoTen"] for p in out["people"]] == ["NGUYỄN VĂN A"]


# ── API ──

def _client(user=USER):
    app = FastAPI()
    app.include_router(rt.router)
    app.add_exception_handler(AppError, app_error_handler)
    app.dependency_overrides[require_auth] = lambda: user
    return TestClient(app)


def test_config_noi_lap_mac_dinh_theo_tai_khoan():
    res = _client().get("/api/v1/assistant/authorization-letter/config").json()
    assert res["place"] == "UBND phường Vũ Ninh, tỉnh Bắc Ninh"
    assert "Chứng thực bản sao từ bản chính" in res["suggestions"]


def _session(monkeypatch, owner="u1", files=2):
    sess = {"_id": "HS-1", "owner_user_id": owner,
            "files": [{"fid": f"f{i}", "name": f"{i}.jpg", "type": "image/jpeg"} for i in range(files)]}
    monkeypatch.setattr(rt.up_store, "get", AsyncMock(return_value=sess))
    monkeypatch.setattr(rt.up_store, "file_to_data_url", lambda sid, f: "data:image/jpeg;base64,AA")
    return sess


def test_extract_doi_chi_so_tep_ra_fid(monkeypatch):
    _session(monkeypatch)
    monkeypatch.setattr(rt, "extract_people", AsyncMock(return_value={
        "people": [{**ex.normalize_person(_person(tepNguon=[1]), 2)}], "unusedFiles": [0], "ocrFailed": []}))
    res = _client().post("/api/v1/assistant/authorization-letter/extract", json={"session_id": "HS-1"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["people"][0]["fids"] == ["f1"] and "tepNguon" not in body["people"][0]
    assert body["unusedFiles"] == ["f0"]


def test_extract_phien_cua_nguoi_khac_bi_chan(monkeypatch):
    _session(monkeypatch, owner="u-khac")
    res = _client().post("/api/v1/assistant/authorization-letter/extract", json={"session_id": "HS-1"})
    assert res.status_code == 403


def test_extract_phien_chua_co_anh(monkeypatch):
    _session(monkeypatch, files=0)
    res = _client().post("/api/v1/assistant/authorization-letter/extract", json={"session_id": "HS-1"})
    assert res.status_code == 400


def test_xoa_phien_khi_dong_tab(monkeypatch):
    _session(monkeypatch)
    deleted = AsyncMock()
    monkeypatch.setattr(rt.up_store, "delete_session", deleted)
    assert _client().delete("/api/v1/assistant/authorization-letter/sessions/HS-1").status_code == 200
    deleted.assert_awaited_once_with("HS-1")


LETTER = {
    "benUyQuyen": {"hoTen": "NGUYỄN VĂN A", "ngaySinh": "01/02/1960", "diaChi": "Thôn Đông",
                   "soDinhDanh": "001060000001", "ngayCap": "10/08/2021", "noiCap": "Bộ Công an"},
    "benDuocUyQuyen": {"hoTen": "TRẦN THỊ B", "ngaySinh": "03/04/1990", "diaChi": "Phường Vũ Ninh",
                       "soDinhDanh": "001190000002"},
    "noiDung": "Thay mặt tôi nộp hồ sơ.", "lapTai": "UBND phường Vũ Ninh", "ngayLap": "30/09/2026",
}


def test_render_docx():
    res = _client().post("/api/v1/assistant/authorization-letter/render", json={**LETTER, "format": "docx"})
    assert res.status_code == 200
    assert "Giay_uy_quyen_NGUYEN_VAN_A.docx" in res.headers["content-disposition"]
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(res.content)).paragraphs)
    assert "Họ tên: NGUYỄN VĂN A, sinh ngày 01/02/1960" in text
    assert "Họ tên: TRẦN THỊ B\n" in text + "\n", "mặc định không ghi ngày sinh bên được ủy quyền (như mẫu gốc)"
    assert "Hôm nay, ngày 30 tháng 09 năm 2026." in text


def test_render_pdf_mot_trang_co_ngay_sinh_ben_duoc_uy_quyen_khi_chon():
    res = _client().post("/api/v1/assistant/authorization-letter/render",
                         json={**LETTER, "format": "pdf", "ghiNgaySinhDuocUyQuyen": True})
    doc = fitz.open(stream=res.content, filetype="pdf")
    assert res.headers["content-type"] == "application/pdf" and len(doc) == 1
    text = doc[0].get_text().replace("\n", " ")
    assert "TRẦN THỊ B, sinh ngày 03/04/1990" in text


def test_render_giay_trong_in_dong_cham_de_viet_tay():
    res = _client().post("/api/v1/assistant/authorization-letter/render", json={"format": "pdf"})
    assert res.status_code == 200
    assert "Họ tên: ....." in fitz.open(stream=res.content, filetype="pdf")[0].get_text()


# ── Mục trên màn chọn thủ tục ──

def test_muc_giay_uy_quyen_chi_hien_voi_extension_moi():
    old = flow._service_list_card({"client_capabilities": {}})
    new = flow._service_list_card({"client_capabilities": {"supportsAuthorizationLetter": True}})
    assert "counterTools" not in old
    assert new["counterTools"]["items"][0]["key"] == "giay-uy-quyen"


def test_che_do_tieng_mong_an_muc_giay_uy_quyen():
    token = flow._TURN_LANG.set("hmong")
    try:
        card = flow._service_list_card({"client_capabilities": {"supportsAuthorizationLetter": True}})
    finally:
        flow._TURN_LANG.reset(token)
    assert "counterTools" not in card


@pytest.mark.parametrize("raw,expected", [({"supportsAuthorizationLetter": True}, True), ({}, False)])
def test_capability_trong_whitelist(raw, expected):
    from app.channels.handfree.chat.router import _clean_client_capabilities
    assert _clean_client_capabilities(raw)["supportsAuthorizationLetter"] is expected



# ── Nhiều người mỗi bên (đồng ủy quyền / ủy quyền cho nhiều người) ──

def _p(name):
    return {"hoTen": name, "ngaySinh": "01/01/1990", "diaChi": "Phường Vũ Ninh", "soDinhDanh": "001090000001",
            "ngayCap": "10/08/2021", "noiCap": "Bộ Công an"}


def _docx_text(content):
    d = docx.Document(io.BytesIO(content))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    return "\n".join(p.text for p in d.paragraphs), "\n".join(cells)


def test_mot_mot_giu_nguyen_mau_goc_khong_bao_gom():
    res = _client().post("/api/v1/assistant/authorization-letter/render", json={
        "format": "docx", "benUyQuyen": [_p("NGƯỜI A")], "benDuocUyQuyen": [_p("NGƯỜI B")]})
    body, _ = _docx_text(res.content)
    assert "Bao gồm" not in body and "Họ tên: NGƯỜI A, sinh ngày 01/01/1990" in body


def test_ben_tu_hai_nguoi_ghi_bao_gom_va_danh_so():
    res = _client().post("/api/v1/assistant/authorization-letter/render", json={
        "format": "docx", "benUyQuyen": [_p("NGƯỜI A1"), _p("NGƯỜI A2")], "benDuocUyQuyen": [_p("NGƯỜI B")]})
    body, sign = _docx_text(res.content)
    uy_quyen = body[body.index("I. BÊN ỦY QUYỀN"):body.index("II. BÊN ĐƯỢC ỦY QUYỀN")]
    duoc = body[body.index("II. BÊN ĐƯỢC ỦY QUYỀN"):body.index("III.")]
    assert "Bao gồm:" in uy_quyen
    assert "1. Họ tên: NGƯỜI A1" in uy_quyen and "2. Họ tên: NGƯỜI A2" in uy_quyen
    assert "Bao gồm" not in duoc, "bên chỉ một người vẫn giữ mẫu gốc"
    assert all(n in sign for n in ("NGƯỜI A1", "NGƯỜI A2", "NGƯỜI B")), "mỗi người một chỗ ký"


@pytest.mark.parametrize("a,b", [(2, 1), (1, 2), (2, 2), (3, 3)])
def test_pdf_nhieu_nguoi_khong_mat_ten_nguoi_ky(a, b):
    """Bảng chữ ký giáp đáy trang từng bị cắt mất hàng (mất tên người ký) mà không báo lỗi."""
    left = [_p(f"NGƯỜI A{i}") for i in range(a)]
    right = [_p(f"NGƯỜI B{i}") for i in range(b)]
    res = _client().post("/api/v1/assistant/authorization-letter/render",
                         json={"format": "pdf", "benUyQuyen": left, "benDuocUyQuyen": right})
    doc = fitz.open(stream=res.content, filetype="pdf")
    text = "".join(p.get_text() for p in doc)
    for person in left + right:
        assert text.count(person["hoTen"]) >= 2, f"thiếu chỗ ký của {person['hoTen']}"
    if (a, b) in ((2, 1), (1, 2)):
        assert len(doc) == 1, "giấy 3 người vẫn gọn một trang cùng chỗ ký"


def test_render_van_nhan_dang_mot_nguoi_cu():
    res = _client().post("/api/v1/assistant/authorization-letter/render",
                         json={"format": "docx", "benUyQuyen": _p("NGƯỜI A"), "benDuocUyQuyen": _p("NGƯỜI B")})
    assert res.status_code == 200 and "Giay_uy_quyen_NGUOI_A.docx" in res.headers["content-disposition"]
