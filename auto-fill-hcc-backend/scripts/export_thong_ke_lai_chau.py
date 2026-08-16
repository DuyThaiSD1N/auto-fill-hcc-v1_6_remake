"""Xuất Excel thống kê hồ sơ phát sinh theo xã/phường (mỗi xã 1 sheet) — phục vụ báo cáo tỉnh.

Cách dùng (chạy tại thư mục gốc backend, cần .env trỏ đúng Mongo):
    python -m scripts.export_thong_ke_lai_chau
    python -m scripts.export_thong_ke_lai_chau --out /tmp/thong_ke.xlsx
    python -m scripts.export_thong_ke_lai_chau --xa "Tân Phong" --xa "Tả Lèng"
    python -m scripts.export_thong_ke_lai_chau --den "06/08/2026 17:00"   # chốt số đến 17h ngày 06/08

Mỗi sheet: STT | Mã thủ tục | Tên thủ tục | Thủ tục thuộc cấp | Phạm vi hỗ trợ | Số lượng hồ sơ đã tiếp nhận.
Chỉ liệt kê thủ tục CÓ hồ sơ phát sinh (> 0). Số liệu cộng dồn từ trace ĐẦU TIÊN của xã đó đến mốc
--den (giờ VN; bỏ trống = hiện tại). Cách đếm giống dashboard: cùng bộ file, hoặc một bộ là tập
con của bộ kia, được tính là một hồ sơ; lượt tách tab vẫn tính riêng từng tab.
"""
import argparse
import asyncio
import unicodedata
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.db.mongo import connect, get_db
from app.procedures.registry import PROCEDURES
from app.reports.procedure_meta import cap_thu_tuc, ma_thu_tuc, pham_vi_ho_tro
from app.traces.metadata import (
    count_distinct_attachment_sets,
    legacy_dossier_count,
    normalized_attachment_name_set,
)

XA_MAC_DINH = ["Tân Phong", "Tả Lèng", "Đoàn Kết", "Bình Lư"]
_VN_TZ = timezone(timedelta(hours=7))  # giờ VN (UTC+7, không DST) — mọi mốc thời gian nhập/hiển thị theo giờ này

def _fold(s: str) -> str:
    s = unicodedata.normalize("NFD", str(s or "").lower()).replace("đ", "d")
    return " ".join("".join(c for c in s if not unicodedata.combining(c)).split())


def _vn_time(dt: datetime | None) -> str:
    if not isinstance(dt, datetime):
        return "—"
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(_VN_TZ).strftime("%d/%m/%Y")


def _vn_datetime(dt: datetime | None) -> str:
    """Như _vn_time nhưng kèm giờ:phút — dùng cho mốc kết thúc (vd '17:00 06/08/2026')."""
    if not isinstance(dt, datetime):
        return "—"
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(_VN_TZ).strftime("%H:%M %d/%m/%Y")


def _parse_den(value: str | None) -> datetime | None:
    """Mốc KẾT THÚC (giờ VN) -> datetime aware. None/rỗng = đến hiện tại; sai định dạng -> thoát."""
    s = str(value or "").strip()
    if not s:
        return None
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y %H", "%Y-%m-%d %H:%M", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).replace(tzinfo=_VN_TZ)
        except ValueError:
            continue
    raise SystemExit(f"[!] Mốc --den không hợp lệ: {s!r}. Dùng 'DD/MM/YYYY HH:MM' (vd '06/08/2026 17:00').")


async def _match_users(db, xa_folded: str) -> list[dict]:
    """User thuộc xã: khớp fold(users.xa) == tên xã (không phân biệt dấu/hoa thường)."""
    out = []
    for u in await db.users.find({}, {"username": 1, "name": 1, "xa": 1, "tinh": 1}).to_list(1000):
        if _fold(u.get("xa")) == xa_folded:
            out.append(u)
    return out


def _dem_ho_so(docs: list[dict]) -> dict[str, dict]:
    """Đếm cùng tiêu chí dashboard: bộ file cho lượt thường, dossier_id cho lượt tách tab."""
    split_dossier_ids: dict[str, set[str]] = {}
    non_split_sets: dict[str, set[frozenset[str]]] = {}
    non_split_empty_ids: dict[str, set[str]] = {}
    requests: dict[str, set[str]] = {}
    estimated: dict[str, int] = {}
    label_theo_trace: dict[str, str] = {}
    first_at: dict[str, datetime] = {}
    for d in docs:
        proc = d.get("procedure") or "—"
        request_id = str(d.get("request_id") or d.get("_id") or "—")
        requests.setdefault(proc, set()).add(request_id)
        label_theo_trace.setdefault(proc, d.get("procedure_label") or proc)
        ca = d.get("created_at")
        if isinstance(ca, datetime) and (proc not in first_at or ca < first_at[proc]):
            first_at[proc] = ca
        ids = [str(value) for value in (d.get("dossier_ids") or []) if value]
        exact = int(d.get("stats_version") or 0) >= 2 and bool(ids)
        if d.get("split") is True:
            if exact:
                split_dossier_ids.setdefault(proc, set()).update(ids)
                continue
            count = legacy_dossier_count(
                attachments=d.get("attachments") or [],
                procedure=proc,
                split=True,
            )
            split_dossier_ids.setdefault(proc, set()).update(
                f"{request_id}:legacy:{index}" for index in range(count)
            )
            estimated[proc] = estimated.get(proc, 0) + count
            continue

        file_set = normalized_attachment_name_set(d.get("attachments") or [])
        if file_set:
            non_split_sets.setdefault(proc, set()).add(file_set)
        else:
            empty_id = ids[0] if exact else request_id
            non_split_empty_ids.setdefault(proc, set()).add(empty_id)

    out: dict[str, dict] = {}
    all_procedures = set(split_dossier_ids) | set(non_split_sets) | set(non_split_empty_ids)
    for proc in all_procedures:
        file_sets = list(non_split_sets.get(proc, set()))
        file_sets.extend(frozenset() for _ in non_split_empty_ids.get(proc, set()))
        inferred_count = count_distinct_attachment_sets(file_sets)
        count = len(split_dossier_ids.get(proc, set())) + inferred_count
        out[proc] = {
            "count": count,
            "requests": len(requests.get(proc, set())),
            "estimated": estimated.get(proc, 0) + inferred_count,
            "label": label_theo_trace.get(proc, proc),
            "first_at": first_at.get(proc),
        }
    return out


# ===== Excel =====
_HEADERS = ["STT", "Mã thủ tục", "Tên thủ tục", "Thủ tục thuộc cấp",
            "Phạm vi hỗ trợ", "Số lượng hồ sơ đã tiếp nhận"]
_WIDTHS = [6, 13, 62, 17, 34, 15]
_THIN = Side(style="thin", color="999999")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")


def _ghi_sheet(wb: Workbook, ten_xa: str, ten_don_vi: str, rows: list[dict],
               tu_ngay: str, den_label: str) -> None:
    ws = wb.create_sheet(title=ten_xa[:31])
    for i, w in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = f"THỐNG KÊ HỒ SƠ TIẾP NHẬN QUA TRỢ LÝ HỖ TRỢ THỦ TỤC HÀNH CHÍNH — {ten_don_vi.upper()}"
    c.font = Font(bold=True, size=13)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:F2")
    c = ws["A2"]
    c.value = f"Số liệu cộng dồn từ lần ghi nhận đầu tiên ({tu_ngay}) đến {den_label}"
    c.font = Font(italic=True, size=10)
    c.alignment = Alignment(horizontal="center")

    hr = 4
    for col, title in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=hr, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[hr].height = 30
    ws.freeze_panes = f"A{hr + 1}"

    r = hr
    for i, row in enumerate(rows, start=1):
        r = hr + i
        values = [i, row["ma"], row["ten"], row["cap"], row["pham_vi"], row["count"]]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.border = _BORDER
            if col in (1, 2, 4, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    cell = ws.cell(row=r, column=1, value="TỔNG CỘNG")
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    total_cell = ws.cell(row=r, column=6, value=sum(x["count"] for x in rows))
    total_cell.font = Font(bold=True)
    total_cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(1, 7):
        ws.cell(row=r, column=col).border = _BORDER
        ws.cell(row=r, column=col).fill = _TOTAL_FILL

    if not rows:
        ws.merge_cells(start_row=hr + 1, start_column=1, end_row=hr + 1, end_column=6)
        ws.cell(row=hr + 1, column=1,
                value="Chưa ghi nhận hồ sơ nào cho đơn vị này.").alignment = Alignment(horizontal="center")


async def main(danh_sach_xa: list[str], out_path: str, den: datetime | None = None) -> None:
    connect()
    db = get_db()
    registry = {p["key"]: p for p in PROCEDURES}

    # Mốc kết thúc: chỉ tính trace created_at <= den (giờ VN). None = đến hiện tại.
    trace_query: dict = {}
    if den is not None:
        trace_query["created_at"] = {"$lte": den}
    den_label = _vn_datetime(den) if den is not None else f"hiện tại ({_vn_datetime(datetime.now(timezone.utc))})"

    wb = Workbook()
    wb.remove(wb.active)  # bỏ sheet mặc định
    thieu_ma: dict[str, str] = {}  # key -> label, để nhắc điền _MA_TT_BO_SUNG

    for ten_xa in danh_sach_xa:
        xa_folded = _fold(ten_xa)
        users = await _match_users(db, xa_folded)
        uids = {str(u["_id"]) for u in users}
        ten_don_vi = next((u.get("name") for u in users if u.get("name")), None) or f"Xã/Phường {ten_xa}"

        # Trace của xã: theo user_id; kèm fallback fold(name/username) chứa tên xã
        # (trace cũ trước khi tài khoản bị tạo lại vẫn được tính).
        docs = []
        cursor = db.traces.find(trace_query, {
            "user_id": 1, "name": 1, "username": 1, "split": 1,
            "request_id": 1, "stats_version": 1, "dossier_ids": 1,
            "procedure": 1, "procedure_label": 1, "attachments": 1, "created_at": 1,
        })
        # Export chạy offline nên stream toàn bộ cursor: không cắt im lặng ở 200.000 trace.
        async for d in cursor:
            if d.get("user_id") in uids or xa_folded in _fold(d.get("name") or d.get("username")):
                docs.append(d)

        theo_thu_tuc = _dem_ho_so(docs)
        first = min((v["first_at"] for v in theo_thu_tuc.values() if v["first_at"]), default=None)

        rows = []
        for key, v in theo_thu_tuc.items():
            if v["count"] <= 0:      # chỉ lấy thủ tục CÓ hồ sơ phát sinh (> 0)
                continue
            entry = registry.get(key)
            ma = ma_thu_tuc(entry)
            if ma == "—":
                thieu_ma[key] = (entry or {}).get("label") or v["label"]
            rows.append({
                "ma": ma,
                "ten": (entry or {}).get("label") or v["label"],
                "cap": cap_thu_tuc(key),
                "pham_vi": pham_vi_ho_tro(entry, key),
                "count": v["count"],
            })
        rows.sort(key=lambda x: (-x["count"], x["ten"]))

        _ghi_sheet(wb, ten_xa, ten_don_vi, rows, _vn_time(first), den_label)
        print(f"[{ten_xa}] tài khoản khớp: {sorted(u.get('username') or '' for u in users) or 'KHÔNG CÓ'}"
              f" · {len(docs)} lượt · {sum(x['count'] for x in rows)} hồ sơ · {len(rows)} thủ tục"
              f" · từ {_vn_time(first)}")

    wb.save(out_path)
    if thieu_ma:
        print("\nCẢNH BÁO — thủ tục chưa có mã TTHC (đang để '—' trong Excel);"
              " điền vào _MA_TT_BO_SUNG rồi chạy lại nếu báo cáo cần mã:")
        for k, label in sorted(thieu_ma.items()):
            print(f"  - {k}: {label[:80]}")
    print(f"\nĐã xuất: {out_path}")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--xa", action="append", default=None,
                   help="Tên xã/phường (lặp lại nhiều lần); mặc định 4 xã Lai Châu")
    p.add_argument("--out", default="thong_ke_lai_chau.xlsx", help="Đường dẫn file Excel xuất ra")
    p.add_argument("--den", default=None,
                   help="Mốc kết thúc (giờ VN), vd '06/08/2026 17:00'. Bỏ trống = đến hiện tại")
    args = p.parse_args()
    asyncio.run(main(args.xa or XA_MAC_DINH, args.out, _parse_den(args.den)))
