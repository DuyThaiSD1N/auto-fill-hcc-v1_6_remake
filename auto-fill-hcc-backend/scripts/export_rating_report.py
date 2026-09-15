"""Xuất Excel BÁO CÁO ĐÁNH GIÁ TRẢI NGHIỆM theo TỈNH (mặc định: Bắc Ninh).

Đọc collection `dossiers` (vòng đời hồ sơ, có cụm `rating` do công dân chấm sau khi nộp).
Gom theo PHƯỜNG → THỦ TỤC, mỗi dòng gồm: số hồ sơ nộp, số đánh giá, tỷ lệ, phân bố 5 mức
(Rất hài lòng…Không hài lòng), số bỏ qua, điểm trung bình.

Vì sao đọc thẳng Mongo (không qua API): báo cáo gấp, chạy 1 lần trên server; không thêm route.
Chỉ ĐỌC, không ghi — an toàn với dữ liệu thật.

`province` trong dossiers = nhãn đầy đủ ("Tỉnh Bắc Ninh") do location_for trả; `ward` = tên phường.
"Số hồ sơ nộp" = có `submit_clicked_at` HOẶC đã có `rating` (đánh giá chỉ diễn ra sau khi nộp;
gộp thêm rating để không bỏ sót hồ sơ mà tín hiệu bấm-nút bị lỡ). "Số đánh giá" = rating có mức 1..5.

Chạy trên server (thư mục backend, dùng venv):
  .venv/bin/python scripts/export_rating_report.py
  .venv/bin/python scripts/export_rating_report.py --tinh "Bắc Ninh" --tu 2026-09-11 --den 2026-09-15
  .venv/bin/python scripts/export_rating_report.py --tinh bacninh --ra bao_cao_bac_ninh.xlsx
"""
import argparse
import os
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pymongo import MongoClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import PieChart, Reference  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from app.config import settings  # noqa: E402
from app.dossiers.rating_card import RATING_CARD  # noqa: E402
from app.locations.lookup import PROVINCES  # noqa: E402

VN_TZ = timezone(timedelta(hours=7))  # mốc thời gian lưu UTC; "ngày" tính theo giờ VN (+7)

# 5 → 1 (Rất hài lòng … Không hài lòng), lấy thẳng từ nguồn nhãn dùng chung 2 kênh.
LEVELS = [item["value"] for item in RATING_CARD["scale"]]          # [5,4,3,2,1]
LEVEL_LABEL = {item["value"]: item["label"] for item in RATING_CARD["scale"]}

# ── Bảng màu báo cáo (tông teal của thương hiệu Trợ lý) ──────────────────────────────────
C_TITLE = PatternFill("solid", fgColor="0E7C66")
C_HEAD = PatternFill("solid", fgColor="127D71")
C_WARD = PatternFill("solid", fgColor="D8EFEA")   # dòng cộng theo phường
C_TOTAL = PatternFill("solid", fgColor="B8E0D6")  # dòng tổng toàn tỉnh
C_ZEBRA = PatternFill("solid", fgColor="F3F9F7")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="C9D6D2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _resolve_province(arg: str) -> tuple[str, str]:
    """Trả (nhãn đầy đủ để query, tên gọn để hiển thị). Nhận slug ('bacninh') hoặc tên."""
    a = (arg or "").strip().lower()
    for p in PROVINCES:
        if a in {p["slug"], p["text"].lower(), p["name"].lower()}:
            return p["text"], p["name"]
    return arg, arg  # không khớp danh mục → dùng thẳng chuỗi làm nhãn query


def _vn_day_start_utc(s: str) -> datetime:
    return datetime.strptime(s, "%Y-%m-%d").replace(tzinfo=VN_TZ).astimezone(timezone.utc)


def _agg_new() -> dict:
    d = {"nop": 0, "dg": 0, "boqua": 0, "tong_diem": 0}
    for lv in LEVELS:
        d[lv] = 0
    return d


def _add(agg: dict, doc: dict) -> None:
    submitted = bool(doc.get("submit_clicked_at")) or bool(doc.get("rating"))
    if submitted:
        agg["nop"] += 1
    rating = doc.get("rating") if isinstance(doc.get("rating"), dict) else None
    if not rating:
        return
    level = rating.get("level")
    if level in LEVEL_LABEL:
        agg["dg"] += 1
        agg[level] += 1
        agg["tong_diem"] += level
    else:  # rating tồn tại nhưng không có mức = công dân bấm "Bỏ qua"
        agg["boqua"] += 1


def _rate(agg: dict) -> str:
    return f"{agg['dg'] / agg['nop'] * 100:.0f}%" if agg["nop"] else "—"


def _avg(agg: dict) -> str:
    return f"{agg['tong_diem'] / agg['dg']:.1f}" if agg["dg"] else "—"


# Cột dữ liệu (sau 2 cột nhãn): Số nộp · Số ĐG · Tỷ lệ · 5 mức · Bỏ qua · Điểm TB
VALUE_HEADERS = ["Số hồ sơ nộp", "Số đánh giá", "Tỷ lệ ĐG"] + \
    [LEVEL_LABEL[lv] for lv in LEVELS] + ["Bỏ qua", "Điểm TB"]


def _value_cells(agg: dict) -> list:
    return [agg["nop"], agg["dg"], _rate(agg)] + [agg[lv] for lv in LEVELS] + [agg["boqua"], _avg(agg)]


def _style_data_row(ws, row: int, ncols: int, fill=None, bold=False):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = BORDER
        cell.alignment = LEFT if c <= 2 else CENTER
        if fill:
            cell.fill = fill
        if bold:
            cell.font = BOLD


def build_stats_sheet(wb, short_name, total, reason_counts, custom_count, lo, hi) -> None:
    """Sheet dashboard: tổng quan + phân bổ mức hài lòng (pie) + phân bổ lý do (pie)."""
    s = wb.create_sheet("Thống kê", 0)  # chèn làm TAB ĐẦU TIÊN

    s.merge_cells("A1:H1")
    t = s.cell(1, 1, f"THỐNG KÊ ĐÁNH GIÁ TRẢI NGHIỆM — {short_name.upper()}")
    t.font = Font(color="FFFFFF", bold=True, size=15)
    t.fill = C_TITLE
    t.alignment = CENTER
    s.row_dimensions[1].height = 30
    s.merge_cells("A2:H2")
    sub = s.cell(2, 1, f"Khoảng dữ liệu: {lo} – {hi}   |   Xuất lúc: {datetime.now(VN_TZ):%H:%M %d/%m/%Y}")
    sub.alignment = CENTER
    sub.font = Font(italic=True, color="35524C")

    def section(row, text):
        s.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        c = s.cell(row, 1, text)
        c.fill = C_HEAD
        c.font = WHITE
        c.alignment = CENTER
        c.border = BORDER

    def kv(row, k, v):
        a, b = s.cell(row, 1, k), s.cell(row, 2, v)
        a.border = b.border = BORDER
        a.alignment = LEFT
        b.alignment = CENTER
        b.font = BOLD

    def header3(row, *hs):
        for c, h in enumerate(hs, start=1):
            cell = s.cell(row, c, h)
            cell.fill = C_WARD
            cell.font = BOLD
            cell.alignment = CENTER
            cell.border = BORDER

    def row3(row, a, b, c):
        s.cell(row, 1, a).alignment = LEFT
        s.cell(row, 2, b).alignment = CENTER
        s.cell(row, 3, c).alignment = CENTER
        for col in (1, 2, 3):
            s.cell(row, col).border = BORDER

    # ── A. Tổng quan ─────────────────────────────────────────────────────────────────────
    section(4, "TỔNG QUAN")
    kv(5, "Số hồ sơ được làm (đã nộp)", total["nop"])
    kv(6, "Số hồ sơ có đánh giá", f"{total['dg']}  ({_rate(total)})")
    kv(7, "Bỏ qua đánh giá", total["boqua"])
    kv(8, "Điểm hài lòng trung bình (/5)", _avg(total))

    # ── B. Phân bổ mức hài lòng + pie ────────────────────────────────────────────────────
    B = 10
    section(B, "PHÂN BỔ MỨC HÀI LÒNG")
    header3(B + 1, "Mức hài lòng", "Số lượng", "Tỷ lệ")
    first = B + 2
    for i, lv in enumerate(LEVELS):
        cnt = total[lv]
        pct = f"{cnt / total['dg'] * 100:.0f}%" if total["dg"] else "—"
        row3(first + i, LEVEL_LABEL[lv], cnt, pct)
    last = first + len(LEVELS) - 1
    if total["dg"]:
        pie = PieChart()
        pie.title = "Phân bổ mức hài lòng"
        pie.add_data(Reference(s, min_col=2, min_row=B + 1, max_row=last), titles_from_data=True)
        pie.set_categories(Reference(s, min_col=1, min_row=first, max_row=last))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        pie.height, pie.width = 8, 13
        s.add_chart(pie, f"E{B}")

    # ── C. Lý do đánh giá (điền nhanh + Khác) + pie ──────────────────────────────────────
    reason_order = list(RATING_CARD["reasonsGood"]) + list(RATING_CARD["reasonsBad"])
    ordered = [x for x in reason_order if reason_counts.get(x)]
    ordered += sorted(x for x in reason_counts if x not in reason_order and reason_counts[x])
    rows = [(x, reason_counts[x]) for x in ordered]
    if custom_count:
        rows.append(("Khác (công dân tự nhập)", custom_count))
    tong_luot = sum(v for _, v in rows)

    C = max(last + 4, B + 17)  # chừa chỗ cho pie B (cao ~16 dòng)
    section(C, "LÝ DO ĐÁNH GIÁ (điền nhanh + Khác)")
    header3(C + 1, "Lý do", "Số lượt chọn", "Tỷ lệ")
    cfirst = C + 2
    for i, (lbl, cnt) in enumerate(rows):
        pct = f"{cnt / tong_luot * 100:.0f}%" if tong_luot else "—"
        row3(cfirst + i, lbl, cnt, pct)
    clast = cfirst + len(rows) - 1
    if rows:
        pie2 = PieChart()
        pie2.title = "Phân bổ lý do (theo lượt chọn)"
        pie2.add_data(Reference(s, min_col=2, min_row=C + 1, max_row=clast), titles_from_data=True)
        pie2.set_categories(Reference(s, min_col=1, min_row=cfirst, max_row=clast))
        pie2.dataLabels = DataLabelList()
        pie2.dataLabels.showPercent = True
        pie2.height, pie2.width = 9, 15
        s.add_chart(pie2, f"E{C}")

    s.column_dimensions["A"].width = 36
    s.column_dimensions["B"].width = 14
    s.column_dimensions["C"].width = 12


def main() -> None:
    ap = argparse.ArgumentParser(description="Xuất Excel báo cáo đánh giá trải nghiệm theo tỉnh.")
    ap.add_argument("--tinh", default="Bắc Ninh", help="Tên hoặc slug tỉnh (mặc định: Bắc Ninh)")
    ap.add_argument("--tu", default="", help="Từ ngày YYYY-MM-DD (giờ VN); bỏ trống = từ đầu")
    ap.add_argument("--den", default="", help="Đến ngày YYYY-MM-DD (giờ VN, bao gồm cả ngày); bỏ trống = đến nay")
    ap.add_argument("--ra", default="", help="Đường dẫn file xuất (.xlsx)")
    args = ap.parse_args()

    query_label, short_name = _resolve_province(args.tinh)

    # Lọc theo tỉnh + CHỈ hồ sơ ĐÃ NỘP + (tùy chọn) khoảng started_at.
    # - Khớp tỉnh linh hoạt: đúng nhãn HOẶC chứa tên gọn (dữ liệu cũ có thể lưu "Bắc Ninh").
    # - "Đã nộp" = có submit_clicked_at HOẶC đã có rating (đánh giá chỉ diễn ra sau khi nộp).
    #   BỎ hồ sơ mới bắt đầu rồi bỏ dở — nếu không, thủ tục chỉ có người mở form sẽ hiện "0 lần làm".
    q: dict = {"$and": [
        {"$or": [
            {"province": query_label},
            {"province": {"$regex": short_name, "$options": "i"}},
        ]},
        {"$or": [
            {"submit_clicked_at": {"$exists": True}},
            {"rating": {"$exists": True}},
        ]},
    ]}
    time_q: dict = {}
    if args.tu:
        time_q["$gte"] = _vn_day_start_utc(args.tu)
    if args.den:
        time_q["$lt"] = _vn_day_start_utc(args.den) + timedelta(days=1)
    if time_q:
        q["$and"].append({"started_at": time_q})

    client = MongoClient(settings.mongo_dsn)
    db = client[settings.mongo_db]
    # Chốt lại "đã nộp" ở tầng docs: query dùng $exists nên lọt cả rating=null (field có mà rỗng);
    # ở đây bỏ hẳn để không sinh dòng thủ tục "0 lần làm".
    docs = [d for d in db.dossiers.find(q)
            if bool(d.get("submit_clicked_at")) or isinstance(d.get("rating"), dict)]

    # Gom: theo (phường, thủ tục), theo phường, theo thủ tục, và tổng tỉnh.
    by_ward_proc: dict = defaultdict(_agg_new)
    by_ward: dict = defaultdict(_agg_new)
    by_proc: dict = defaultdict(_agg_new)
    total = _agg_new()
    reason_counts: dict = defaultdict(int)   # lý do "điền nhanh" (chip) → số lượt chọn
    custom_count = 0                          # số phiếu công dân TỰ NHẬP ý kiến → gộp "Khác"
    dates = []
    for d in docs:
        ward = (d.get("ward") or "").strip() or "(chưa rõ phường)"
        proc = (d.get("procedure_label") or d.get("procedure") or "").strip() or "(chưa rõ thủ tục)"
        for bucket in (by_ward_proc[(ward, proc)], by_ward[ward], by_proc[proc], total):
            _add(bucket, d)
        r = d.get("rating")
        if isinstance(r, dict) and r.get("level") in LEVEL_LABEL:
            for reason in (r.get("reasons") or []):
                lbl = str(reason).strip()
                if lbl:
                    reason_counts[lbl] += 1
            if str(r.get("note") or "").strip():   # có gõ/nói ý kiến riêng = "Khác"
                custom_count += 1
        if d.get("started_at"):
            dates.append(d["started_at"])

    lo = min(dates).astimezone(VN_TZ).strftime("%d/%m/%Y") if dates else "—"
    hi = max(dates).astimezone(VN_TZ).strftime("%d/%m/%Y") if dates else "—"

    wb = Workbook()

    # ══ Sheet 1: Tổng hợp theo phường & thủ tục ══════════════════════════════════════════
    ws = wb.active
    ws.title = "Tổng hợp"
    headers = ["Phường / xã", "Thủ tục"] + VALUE_HEADERS
    ncols = len(headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    t = ws.cell(row=1, column=1, value=f"BÁO CÁO ĐÁNH GIÁ TRẢI NGHIỆM — {short_name.upper()}")
    t.font = Font(color="FFFFFF", bold=True, size=15)
    t.fill = C_TITLE
    t.alignment = CENTER
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    sub = ws.cell(row=2, column=1, value=(
        f"Khoảng dữ liệu: {lo} – {hi}   |   Xuất lúc: {datetime.now(VN_TZ):%H:%M %d/%m/%Y}   |   "
        f"Tổng: {total['nop']} hồ sơ nộp · {total['dg']} đánh giá ({_rate(total)}) · điểm TB {_avg(total)}"
    ))
    sub.alignment = CENTER
    sub.font = Font(italic=True, color="35524C")

    hrow = 4
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hrow, column=c, value=h)
        cell.fill = C_HEAD
        cell.font = WHITE
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[hrow].height = 30

    r = hrow + 1
    zebra = False
    for ward in sorted(by_ward_proc_keys := {k[0] for k in by_ward_proc}):
        procs = sorted(p for (w, p) in by_ward_proc if w == ward)
        for proc in procs:
            agg = by_ward_proc[(ward, proc)]
            ws.cell(row=r, column=1, value=ward)
            ws.cell(row=r, column=2, value=proc)
            for i, v in enumerate(_value_cells(agg), start=3):
                ws.cell(row=r, column=i, value=v)
            _style_data_row(ws, r, ncols, fill=(C_ZEBRA if zebra else None))
            zebra = not zebra
            r += 1
        # Dòng cộng theo phường
        ws.cell(row=r, column=1, value=ward)
        ws.cell(row=r, column=2, value=f"▸ Cộng {ward}")
        for i, v in enumerate(_value_cells(by_ward[ward]), start=3):
            ws.cell(row=r, column=i, value=v)
        _style_data_row(ws, r, ncols, fill=C_WARD, bold=True)
        r += 1
        zebra = False

    # Dòng tổng toàn tỉnh
    ws.cell(row=r, column=1, value="TỔNG")
    ws.cell(row=r, column=2, value=f"TỔNG TOÀN {short_name.upper()}")
    for i, v in enumerate(_value_cells(total), start=3):
        ws.cell(row=r, column=i, value=v)
    _style_data_row(ws, r, ncols, fill=C_TOTAL, bold=True)

    ws.freeze_panes = ws.cell(row=hrow + 1, column=1)
    widths = [22, 40, 12, 11, 9] + [11] * len(LEVELS) + [9, 9]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ══ Sheet 2 & 3: tóm tắt theo phường / theo thủ tục ══════════════════════════════════
    def _summary_sheet(title: str, label_head: str, buckets: dict):
        s = wb.create_sheet(title)
        heads = [label_head] + VALUE_HEADERS
        nc = len(heads)
        s.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc)
        h0 = s.cell(row=1, column=1, value=f"{title.upper()} — {short_name}")
        h0.font = Font(color="FFFFFF", bold=True, size=13)
        h0.fill = C_TITLE
        h0.alignment = CENTER
        s.row_dimensions[1].height = 26
        for c, h in enumerate(heads, start=1):
            cell = s.cell(row=3, column=c, value=h)
            cell.fill = C_HEAD
            cell.font = WHITE
            cell.alignment = CENTER
            cell.border = BORDER
        s.row_dimensions[3].height = 28
        rr = 4
        # Sắp theo số hồ sơ nộp giảm dần cho dễ đọc.
        for key in sorted(buckets, key=lambda k: buckets[k]["nop"], reverse=True):
            s.cell(row=rr, column=1, value=key)
            for i, v in enumerate(_value_cells(buckets[key]), start=2):
                s.cell(row=rr, column=i, value=v)
            _style_data_row(s, rr, nc, fill=(C_ZEBRA if rr % 2 else None))
            rr += 1
        s.cell(row=rr, column=1, value="TỔNG")
        for i, v in enumerate(_value_cells(total), start=2):
            s.cell(row=rr, column=i, value=v)
        _style_data_row(s, rr, nc, fill=C_TOTAL, bold=True)
        s.freeze_panes = s.cell(row=4, column=1)
        for i, w in enumerate([34, 12, 11, 9] + [11] * len(LEVELS) + [9, 9], start=1):
            s.column_dimensions[get_column_letter(i)].width = w

    _summary_sheet("Theo phường", "Phường / xã", by_ward)
    _summary_sheet("Theo thủ tục", "Thủ tục", by_proc)
    build_stats_sheet(wb, short_name, total, reason_counts, custom_count, lo, hi)  # tab đầu

    out = args.ra or f"bao_cao_danh_gia_{_resolve_province(args.tinh)[1].lower().replace(' ', '_')}_{datetime.now(VN_TZ):%Y%m%d}.xlsx"
    wb.save(out)
    print(f"✅ Đã xuất: {out}")
    print(f"   Tỉnh: {short_name} | Khoảng: {lo}–{hi} | {len(docs)} hồ sơ")
    print(f"   Tổng nộp {total['nop']} · đánh giá {total['dg']} ({_rate(total)}) · bỏ qua {total['boqua']} · điểm TB {_avg(total)}")
    if not docs:
        print("   ⚠️  KHÔNG có hồ sơ nào khớp — kiểm tra lại tên tỉnh (--tinh) hoặc khoảng ngày.")


if __name__ == "__main__":
    main()
