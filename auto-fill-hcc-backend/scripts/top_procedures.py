"""Xuất EXCEL: top thủ tục nhiều HỒ SƠ THỰC TẾ phát sinh nhất (mặc định top 20).

Cột: Mã thủ tục · Tên thủ tục · Số hồ sơ phát sinh.

"Hồ sơ phát sinh (thực tế)" = hồ sơ riêng biệt sau khi gộp trùng — dùng ĐÚNG cách đếm của
dashboard (repo.stats: gộp các lượt bấm cùng một bộ giấy tờ, KHÔNG tính lượt bấm/số request,
không đếm trùng bước đính kèm). "Mã thủ tục" = mã TTHC (maThuTuc) lấy từ registry; thủ tục
không mang mã trên URL thì để trống, cột phụ ghi khóa nội bộ để đối chiếu.

Chạy trên server (trong thư mục auto-fill-hcc-backend):
  python scripts/top_procedures.py                       # top 20, toàn thời gian, mọi tài khoản
  python scripts/top_procedures.py --top 30
  python scripts/top_procedures.py --scope official      # chỉ tài khoản phường (loại admin/test)
  python scripts/top_procedures.py --from 2026-08-01 --to 2026-08-20
  python scripts/top_procedures.py --out top_thu_tuc.xlsx

Trong Docker:
  docker compose -f compose.prod.yml exec app python scripts/top_procedures.py
  docker compose -f compose.prod.yml cp app:/app/top_thu_tuc_YYYY-MM-DD.xlsx ./   # lấy file về
"""
import argparse
import asyncio
import os
import re
import sys
from datetime import datetime, timedelta, timezone

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.db.mongo import close, connect  # noqa: E402
from app.procedures.registry import get_procedure  # noqa: E402
from app.traces import repo  # noqa: E402

VN_TZ = timezone(timedelta(hours=7))  # created_at lưu UTC; "ngày" nhập vào tính theo giờ VN
_MA_RE = re.compile(r"maThuTuc=([0-9.]+)")


def _parse_day(value: str | None, *, end: bool) -> datetime | None:
    """Đầu/cuối ngày (giờ VN) → UTC. date_to là mốc LOẠI TRỪ (nửa mở) đúng như repo.stats."""
    if not value:
        return None
    day = datetime.strptime(value, "%Y-%m-%d").replace(tzinfo=VN_TZ)
    if end:
        day += timedelta(days=1)
    return day.astimezone(timezone.utc)


def _ma_thu_tuc(key: str) -> str:
    """Mã TTHC (maThuTuc) lấy từ detect.urlIncludes của registry; không có thì trả rỗng."""
    proc = get_procedure(key) or {}
    for item in (proc.get("detect") or {}).get("urlIncludes") or []:
        m = _MA_RE.search(str(item))
        if m:
            return m.group(1)
    return ""


async def main() -> None:
    ap = argparse.ArgumentParser(description="Xuất Excel top thủ tục nhiều hồ sơ thực tế nhất")
    ap.add_argument("--top", type=int, default=20, help="Số thủ tục (mặc định 20)")
    ap.add_argument("--scope", choices=["all", "official"], default="all",
                    help="all = mọi tài khoản; official = chỉ phường (commune/province)")
    ap.add_argument("--from", dest="date_from", help="Từ ngày YYYY-MM-DD (giờ VN)")
    ap.add_argument("--to", dest="date_to", help="Đến ngày YYYY-MM-DD (giờ VN, tính cả ngày)")
    ap.add_argument("--out", help="Đường dẫn file Excel (mặc định top_thu_tuc_<ngày>.xlsx)")
    args = ap.parse_args()

    date_from = _parse_day(args.date_from, end=False)
    date_to = _parse_day(args.date_to, end=True)

    connect()
    try:
        res = await repo.stats(date_from=date_from, date_to=date_to, scope=args.scope)
    finally:
        close()

    procedures = res.get("procedures") or []
    top = procedures[: max(args.top, 1)]

    span = "toan-thoi-gian"
    if args.date_from or args.date_to:
        span = f"{args.date_from or 'dau'}_{args.date_to or 'nay'}"
    out_path = args.out or f"top_thu_tuc_{datetime.now(VN_TZ):%Y-%m-%d}.xlsx"

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Top thủ tục"

    headers = ["STT", "Mã thủ tục", "Tên thủ tục", "Số hồ sơ phát sinh", "Khóa nội bộ"]
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(color="FFFFFF", bold=True)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, p in enumerate(top, 1):
        key = p.get("key") or ""
        ws.cell(i + 1, 1, i)
        ws.cell(i + 1, 2, _ma_thu_tuc(key))
        ws.cell(i + 1, 3, p.get("label") or key or "—")
        ws.cell(i + 1, 4, int(p.get("count") or 0))
        ws.cell(i + 1, 5, key)

    for col, w in zip("ABCDE", [6, 16, 60, 20, 32]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    wb.save(out_path)

    # Tóm tắt ra màn hình để chạy xong biết ngay.
    print(f"Phạm vi: {args.scope} | Thời gian: {span} | Tổng hồ sơ thực tế: {res.get('totalDossiers', 0):,}")
    print(f"Đã ghi Excel ({len(top)} thủ tục): {os.path.abspath(out_path)}")
    for i, p in enumerate(top[:10], 1):
        print(f"  {i:>2}. {int(p.get('count') or 0):>6,}  {p.get('label') or p.get('key')}")
    if len(top) > 10:
        print(f"  … và {len(top) - 10} thủ tục nữa (xem file Excel).")


if __name__ == "__main__":
    asyncio.run(main())
